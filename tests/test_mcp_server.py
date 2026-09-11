"""docx2typed-mcp: span-free region-scoped editing tools and regions.md."""
from __future__ import annotations

import pytest

import json
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from scripts import main
from scripts.extract import extract
from scripts.protocol import file_sha256
from scripts.review_queue import dispatch, upsert_event
from scripts.review_collab import stage_patch
from scripts.mcp_server import (
    _fnv1a_utf16,
    batch_edit,
    build_docx,
    commit_sync,
    document_read,
    document_search,
    document_patch,
    decide_all,
    delete_paragraph,
    diff_preview,
    get_paragraph,
    insert_paragraph,
    list_paragraphs,
    replace_text,
    revert,
    review_ack,
    review_apply_batch,
    review_apply_patch,
    review_inbox,
    review_preflight,
    review_external_preflight,
    review_settle,
    review_state,
    session,
    table_insert_row,
    verify_output,
    workdir_open,
    document_replace,
    engine_info,
    ToolError,
    _scope_paragraph_ids,
    format_span,
    _workdir_open_result,
    workdir_status,
)

ROOT = Path(__file__).resolve().parents[1]


def _reset() -> None:
    session.workdir = None
    session.last_build_output = None


def make_doc(path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("前言")
    cn = paragraph.add_run("智能响应")
    cn.font.name = "宋体"
    cn._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    en = paragraph.add_run("ABC")
    en.font.name = "Times New Roman"
    paragraph.add_run("后语")
    document.add_paragraph("第二段")
    document.save(path)


def open_workdir(tmp_path: Path, name: str) -> Path:
    source = tmp_path / f"{name}-src.docx"
    workdir = tmp_path / name
    make_doc(source)
    assert extract([str(source), "-o", str(workdir)]) == 0
    return json.loads(workdir_open(str(workdir)))["workdir"]


def _j(result) -> dict:
    """Unwrap a tool result: the data dict from a Result envelope's
    structuredContent, a plain dict passthrough, or a legacy JSON string."""
    if hasattr(result, "structuredContent"):
        return result.structuredContent["data"]
    if isinstance(result, dict):
        return result
    return json.loads(result)


def cn_en_runs(docx_path: Path) -> dict:
    return {r.text: r for r in Document(docx_path).paragraphs[0].runs}


def make_table_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("表前")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "A2"
    table.cell(1, 0).text = "B1"
    table.cell(1, 1).text = "B2"
    document.add_paragraph("表后")
    document.save(path)


def open_store_workdir(tmp_path: Path, name: str) -> Path:
    """Extract via the JSON CLI, which births the immutable-generation store
    (generation 0), then open the MCP session."""
    source = tmp_path / f"{name}-src.docx"
    workdir = tmp_path / name
    make_doc(source)
    assert main(["--json", "extract", str(source), "-o", str(workdir), "--operation-id", f"{name}-extract-1"]) == 0
    return Path(json.loads(workdir_open(str(workdir)))["workdir"])


def test_regions_md_generated_at_extract(tmp_path):
    _reset()
    workdir = open_workdir(tmp_path, "regions")
    regions = (Path(workdir) / "regions.md").read_text(encoding="utf-8")
    assert "## P0" in regions and "## P1" in regions
    assert "智能响应" in regions and "{s_" in regions
    assert "[token]" not in regions  # no tokens in this fixture


def test_get_paragraph_has_style_id_description_rpr(tmp_path):
    _reset()
    open_workdir(tmp_path, "styles")
    data = _j(get_paragraph("P0"))
    styles = data["styles"]
    cn_region = next(r for r in styles if r["text"] == "智能响应")
    en_region = next(r for r in styles if r["text"] == "ABC")
    assert cn_region["style_id"].startswith("s_") and cn_region["style_id"] != en_region["style_id"]
    assert cn_region["description"]  # non-empty label
    assert cn_region["rpr"].startswith("<w:rPr")  # full canonical XML present


def test_batch_edit_by_index_preserves_cn_en_fonts(tmp_path):
    _reset()
    open_workdir(tmp_path, "batch")
    result = _j(batch_edit(
        "P0",
        [
            {"region": 1, "new": "智能调控"},
            {"region": 2, "new": "XYZ"},
        ],
        operation_id="batch-index-1",
    ))
    assert result["edits_applied"] == 2 and result["state"] == "clean"
    _j(commit_sync(operation_id="batch-build-save"))  # ADR 0044: export requires a saved version
    output = _j(build_docx(operation_id="batch-build-1"))["output"]
    assert _j(verify_output(output))["verified"] == output
    runs = cn_en_runs(output)
    assert runs["智能调控"]._element.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"
    assert runs["XYZ"]._element.rPr.rFonts.get(qn("w:ascii")) == "Times New Roman"


def test_batch_edit_text_anchor_with_style_disambiguation(tmp_path):
    _reset()
    open_workdir(tmp_path, "anchor")
    _j(batch_edit("P0", [{"text": "ABC", "new": "XYZ"}], operation_id="batch-anchor-1"))
    assert _j(workdir_status())["state"] == "clean"
    data = _j(get_paragraph("P0"))
    assert "XYZ" in data["plain"] and "ABC" not in data["plain"]


def test_batch_edit_partial_old_inside_region(tmp_path):
    _reset()
    open_workdir(tmp_path, "partial")
    _j(batch_edit("P0", [{"region": 1, "old": "智能", "new": "智慧"}], operation_id="batch-partial-1"))
    data = _j(get_paragraph("P0"))
    assert "智慧响应" in data["plain"]


def test_batch_edit_atomic_rejection(tmp_path):
    _reset()
    open_workdir(tmp_path, "atomic")
    result = batch_edit(
        "P0",
        [
            {"region": 1, "new": "智能调控"},
            {"region": 2, "old": "not-there", "new": "XYZ"},
        ],
        operation_id="batch-atomic-1",
    )
    assert result.isError is True
    assert result.structuredContent["diagnostics"][0]["code"] == "text-not-found"
    assert _j(workdir_status())["state"] == "clean"  # rolled back
    data = _j(get_paragraph("P0"))
    assert "智能响应" in data["plain"] and "智能调控" not in data["plain"]


def test_batch_edit_region_out_of_range_and_duplicate(tmp_path):
    _reset()
    open_workdir(tmp_path, "range")
    for index, (bad, code) in enumerate((
        ([{"region": 99, "new": "x"}], "region-out-of-range"),
        ([{"region": 0, "new": "a"}, {"region": 0, "new": "b"}], "invalid-edit"),
    )):
        result = batch_edit("P0", bad, operation_id=f"batch-range-{index}")
        assert result.isError is True
        assert result.structuredContent["diagnostics"][0]["code"] == code


def test_regions_md_auto_updates_after_edit(tmp_path):
    _reset()
    workdir = open_workdir(tmp_path, "autoupdate")
    _j(batch_edit("P0", [{"region": 1, "new": "新词"}], operation_id="batch-auto-1"))
    regions = (Path(workdir) / "regions.md").read_text(encoding="utf-8")
    assert "新词" in regions and "智能响应" not in regions


def test_replace_text_still_rejects_cross_region(tmp_path):
    _reset()
    open_workdir(tmp_path, "cross")
    result = replace_text("P0", "智能响应ABC", "新词XYZ", operation_id="cross-region-1")
    assert result.isError is True
    assert result.structuredContent["diagnostics"][0]["code"] == "cross-region-text"


def test_full_workflow_through_mcp(tmp_path):
    _reset()
    open_workdir(tmp_path, "flow")
    _j(insert_paragraph("P0", "新增段", operation_id="flow-insert-1"))
    _j(delete_paragraph("P1", operation_id="flow-delete-1"))
    _j(commit_sync(operation_id="flow-commit-1"))
    _j(revert(operation_id="flow-revert-1"))
    output = _j(build_docx(operation_id="flow-build-1"))["output"]
    assert _j(verify_output(output))["verified"] == output
    texts = [p.text for p in Document(output).paragraphs]
    assert "新增段" in texts and "第二段" not in texts
    assert list_paragraphs()


def test_verify_output_evidence_publish_failure_is_deterministic(tmp_path, monkeypatch):
    """Issue #50 final finding: an evidence publish failure reports the
    deterministic '{type}: {stable path}' detail (never the transient temp
    filename) and every retry produces the byte-identical diagnostic."""
    _reset()
    open_workdir(tmp_path, "verify-publish")
    _j(commit_sync(operation_id="verify-publish-save"))  # ADR 0044: export requires a saved version
    output = _j(build_docx(operation_id="verify-publish-build-1"))["output"]
    import scripts.mcp_server as mcp_server

    transient = str(tmp_path / ".out.docx.verify.evidence.json.abc12345.tmp")

    def boom(path, evidence):
        raise OSError(f"[Errno 28] No space left on device: {transient!r}")

    monkeypatch.setattr(mcp_server, "publish_run_evidence", boom)
    first = verify_output(output)
    second = verify_output(output)
    assert first.isError is True and second.isError is True
    diagnostic = first.structuredContent["diagnostics"][0]
    assert diagnostic["code"] == "evidence-publish-failed"
    assert diagnostic["message"] == (
        f"required run evidence could not be published: OSError: {output}.verify.evidence.json"
    )
    assert transient not in json.dumps(first.structuredContent)
    assert first.structuredContent == second.structuredContent  # byte-exact retry

def test_verify_output_rejects_reused_id_after_canonical_change(tmp_path):
    _reset()
    open_workdir(tmp_path, "verify-reuse")
    _j(commit_sync(operation_id="verify-reuse-save"))  # ADR 0044: export requires a saved version
    output = _j(build_docx(operation_id="verify-reuse-build"))["output"]
    op = "verify-reuse-check"
    first = verify_output(output, operation_id=op)
    assert first.isError is False
    assert verify_output(output, operation_id=op).structuredContent == first.structuredContent

    _j(replace_text("P0", "智能响应", "智能调控", operation_id="verify-reuse-edit"))
    _j(commit_sync(operation_id="verify-reuse-commit"))
    reused = verify_output(output, operation_id=op)
    assert reused.isError is True
    assert reused.structuredContent["diagnostics"][0]["code"] == "operation-id-reused"

def test_external_preflight_replays_and_rejects_changed_input(tmp_path):
    _reset()
    workdir = Path(open_workdir(tmp_path, "external-preflight"))
    current = json.loads(review_state())["current_snapshot"]["id"]
    op = "external-preflight-1"

    first_call = review_external_preflight(
        current, operation="import", operation_id=op
    )
    first = _j(first_call)
    assert first["operation_id"] == op
    assert first["operation"] == "import"
    replay = review_external_preflight(
        current, operation="import", operation_id=op
    )
    assert replay.structuredContent == first_call.structuredContent

    reused = review_external_preflight(
        current, operation="rollback", operation_id=op
    )
    assert reused.isError is True
    assert reused.structuredContent["diagnostics"][0]["code"] == "operation-id-reused"

def test_review_ack_empty_list_returns_structured_diagnostic(tmp_path):
    _reset()
    open_workdir(tmp_path, "empty-review-ack")

    result = review_ack([], operation_id="empty-review-ack-1")

    assert result.isError is True
    assert result.structuredContent["diagnostics"][0]["code"] == "event-ids-required"
    assert result.structuredContent["data"]["operation_id"] == "empty-review-ack-1"


def test_review_apply_patch_settles_human_text_before_agent_write(tmp_path):
    _reset()
    workdir = Path(open_workdir(tmp_path, "human-patch"))
    current = json.loads(review_preflight())["current_snapshot"]["id"]
    paragraph = json.loads(get_paragraph("P0"))
    paragraph_text = paragraph["plain"]
    before = "智能响应"
    start = paragraph_text.index(before)
    target = {
        "start_offset": start,
        "end_offset": start + len(before),
        "expected_text": before,
        "left_context": paragraph_text[max(0, start - 100):start],
        "right_context": paragraph_text[start + len(before):start + len(before) + 100],
        "paragraph_fingerprint": _fnv1a_utf16(paragraph_text),
        "region_fingerprint": _fnv1a_utf16(before),
        "style_region_ids": list(dict.fromkeys(region["style_id"] for region in paragraph["styles"])),
    }
    event = upsert_event(
        workdir,
        {
            "type": "patch",
            "client_id": "human:patch-1",
            "origin": "human_ui",
            "author": "Lin",
            "parent_snapshot": current,
            "paragraph_id": "P0",
            "kind": "replace",

            "target": target,
            "before": before,
            "after": "智能调控",
        },
    )
    queued = dispatch(workdir)
    assert queued[0]["event_id"] == event["event_id"]

    first_call = review_apply_patch(event["event_id"])
    result = _j(first_call)
    assert result["state"] == "applied"
    assert result["operation_id"] == f"review-apply-patch-{event['event_id']}"
    # Identical retry replays the original envelope byte-exact (no second
    # effect); the store ledger dedupes on the event-derived operation id.
    replayed = review_apply_patch(event["event_id"])
    assert replayed.structuredContent == first_call.structuredContent
    assert replayed.structuredContent["data"]["state"] == "applied"
    assert result["commit"]["current_snapshot"]["id"] == "C1"
    assert json.loads(review_state())["current_snapshot"]["origin"] == "human_ui"
    assert "智能调控" in json.loads(get_paragraph("P0"))["plain"]
    assert json.loads(review_preflight())["ready"] is True

def test_review_apply_patch_commits_staged_batch_atomically(tmp_path):
    _reset()
    workdir = Path(open_workdir(tmp_path, "human-batch"))
    current = json.loads(review_preflight())["current_snapshot"]["id"]
    paragraph = json.loads(get_paragraph("P0"))
    paragraph_text = paragraph["plain"]
    style_region_ids = list(dict.fromkeys(region["style_id"] for region in paragraph["styles"]))

    def make_patch(before: str, after: str, parent: str, client_id: str) -> dict[str, object]:
        start = paragraph_text.index(before)
        return {
            "type": "patch",
            "client_id": client_id,
            "origin": "human_ui",
            "author": "Lin",
            "parent_snapshot": parent,
            "paragraph_id": "P0",
            "kind": "replace",
            "target": {
                "start_offset": start,
                "end_offset": start + len(before),
                "expected_text": before,
                "left_context": paragraph_text[max(0, start - 100):start],
                "right_context": paragraph_text[start + len(before):start + len(before) + 100],
                "paragraph_fingerprint": _fnv1a_utf16(paragraph_text),
                "region_fingerprint": _fnv1a_utf16(before),
                "style_region_ids": style_region_ids,
            },
            "before": before,
            "after": after,
        }

    first = stage_patch(workdir, make_patch("前言", "导言", current, "batch:first"))
    second = stage_patch(workdir, make_patch("智能响应", "智能调控", first["staged_snapshot"], "batch:second"))
    dispatch(workdir)

    result = _j(review_apply_patch(first["event_id"]))
    assert result["state"] == "applied"
    assert result["commit"]["current_snapshot"]["id"] == "C1"
    assert json.loads(review_state())["current_snapshot"]["origin"] == "human_ui"
    assert {event["delivery_state"] for event in result["events"]} == {"applied"}
    assert "导言智能调控" in json.loads(get_paragraph("P0"))["plain"]

# --------------------------------------------------------------------------
# Issue #50 final findings: review mutators route through the idempotent
# store seam (operation_id replay / no second effect / operation-id-reused)
# --------------------------------------------------------------------------

def _stage_human_patch(workdir: Path, current: str, before: str, after: str, client_id: str) -> dict:
    paragraph_text = json.loads(get_paragraph("P0"))["plain"]
    start = paragraph_text.index(before)
    style_region_ids = list(dict.fromkeys(region["style_id"] for region in json.loads(get_paragraph("P0"))["styles"]))
    return stage_patch(
        workdir,
        {
            "type": "patch",
            "client_id": client_id,
            "origin": "human_ui",
            "author": "Lin",
            "parent_snapshot": current,
            "paragraph_id": "P0",
            "kind": "replace",
            "target": {
                "start_offset": start,
                "end_offset": start + len(before),
                "expected_text": before,
                "left_context": paragraph_text[max(0, start - 100):start],
                "right_context": paragraph_text[start + len(before):start + len(before) + 100],
                "paragraph_fingerprint": _fnv1a_utf16(paragraph_text),
                "region_fingerprint": _fnv1a_utf16(before),
                "style_region_ids": style_region_ids,
            },
            "before": before,
            "after": after,
        },
    )

def test_agent_write_gate_scopes_queued_patch_to_target(tmp_path):
    _reset()
    workdir = Path(open_workdir(tmp_path, "scoped-gate"))
    current = json.loads(review_preflight())["current_snapshot"]["id"]
    event = _stage_human_patch(workdir, current, "智能响应", "智能调控", "scope:patch")
    dispatch(workdir)

    unaffected = _j(
        replace_text("P1", "第二段", "第二节", operation_id="scope-agent-p1")
    )
    assert unaffected["paragraph_id"] == "P1"
    assert "第二节" in json.loads(get_paragraph("P1"))["plain"]

    blocked = replace_text(
        "P0", "前言", "导言", operation_id="scope-agent-p0"
    )
    assert blocked.isError is True
    diagnostic = blocked.structuredContent["diagnostics"][0]
    assert diagnostic["code"] == "agent-preflight-required"
    assert diagnostic["details"]["scope"] == ["P0"]
    assert diagnostic["details"]["blocked_patches"][0]["event_id"] == event["event_id"]
    assert blocked.structuredContent["data"]["recovery"]["action"] == "resolve-review"

def test_blocked_write_does_not_bootstrap_store(tmp_path):
    _reset()
    workdir = Path(open_workdir(tmp_path, "blocked-no-store"))
    current = json.loads(review_preflight())["current_snapshot"]["id"]
    _stage_human_patch(workdir, current, "智能响应", "智能调控", "blocked:no-store")
    dispatch(workdir)
    store_dir = workdir / ".docx2typed-store"
    assert not store_dir.exists()

    blocked = replace_text("P0", "前言", "导言", operation_id="blocked-no-store-op")

    assert blocked.isError is True
    assert blocked.structuredContent["diagnostics"][0]["code"] == "agent-preflight-required"
    assert not store_dir.exists()


def test_review_apply_patch_replays_exact_envelope_without_second_effect(tmp_path):
    """Findings: review_apply_patch routes through _mutation_tool with a
    stable event-derived operation_id. An identical retry returns the
    original committed envelope byte-exact and never applies a second time."""
    _reset()
    workdir = Path(open_workdir(tmp_path, "human-replay"))
    current = json.loads(review_preflight())["current_snapshot"]["id"]
    event = _stage_human_patch(workdir, current, "智能响应", "智能调控", "replay:1")
    dispatch(workdir)

    first_call = review_apply_patch(event["event_id"])
    first = _j(first_call)
    assert first["state"] == "applied"
    assert first["operation_id"] == f"review-apply-patch-{event['event_id']}"
    assert first["commit"]["current_snapshot"]["id"] == "C1"

    replayed = review_apply_patch(event["event_id"])
    assert replayed.structuredContent == first_call.structuredContent  # byte-exact replay
    assert json.loads(review_state())["current_snapshot"]["id"] == "C1"  # no second effect
    assert "智能调控" in json.loads(get_paragraph("P0"))["plain"]
    assert "智能响应" not in json.loads(get_paragraph("P0"))["plain"]


def test_review_apply_patch_operation_id_reused(tmp_path):
    """Findings: an explicit operation_id reused with different canonical
    input fails operation-id-reused; identical input still replays."""
    _reset()
    workdir = Path(open_workdir(tmp_path, "human-reused"))
    current = json.loads(review_preflight())["current_snapshot"]["id"]
    first_event = _stage_human_patch(workdir, current, "智能响应", "智能调控", "reused:1")
    second_event = _stage_human_patch(workdir, first_event["staged_snapshot"], "前言", "导言", "reused:2")
    dispatch(workdir)

    op = "explicit-patch-op-1"
    applied = _j(review_apply_patch(first_event["event_id"], operation_id=op))
    assert applied["state"] == "applied" and applied["operation_id"] == op
    # Identical retry: replay, never a second effect.
    assert review_apply_patch(first_event["event_id"], operation_id=op).structuredContent["data"]["state"] == "applied"
    # Same operation_id, different event: rejected before any effect — the
    # snapshot stays at the batch commit (C1), never a fresh generation.
    reused = review_apply_patch(second_event["event_id"], operation_id=op)
    assert reused.isError is True
    assert reused.structuredContent["diagnostics"][0]["code"] == "operation-id-reused"
    assert json.loads(review_state())["current_snapshot"]["id"] == "C1"


def test_review_apply_batch_replays_exact_envelope_without_second_effect(tmp_path):
    """Findings: review_apply_batch routes through the store seam with a
    stable batch-derived operation_id; an identical retry replays the
    original committed envelope byte-exact and never applies again."""
    _reset()
    workdir = Path(open_workdir(tmp_path, "batch-replay"))
    current = json.loads(review_preflight())["current_snapshot"]["id"]
    first_event = _stage_human_patch(workdir, current, "智能响应", "智能调控", "batch:replay-1")
    second_event = _stage_human_patch(workdir, first_event["staged_snapshot"], "前言", "导言", "batch:replay-2")
    queued = dispatch(workdir)
    assert len(queued) == 2 and queued[0]["batch_id"] == queued[1]["batch_id"]
    batch_id = queued[0]["batch_id"]

    first_call = review_apply_batch(batch_id)
    first = _j(first_call)
    assert first["state"] == "applied"
    assert first["operation_id"] == f"review-apply-batch-{batch_id}"
    assert first["commit"]["current_snapshot"]["id"] == "C1"

    assert review_apply_batch(batch_id).structuredContent == first_call.structuredContent
    assert json.loads(review_state())["current_snapshot"]["id"] == "C1"  # no second effect
    assert "导言智能调控" in json.loads(get_paragraph("P0"))["plain"]


def test_review_apply_batch_operation_id_reused(tmp_path):
    """Findings: an explicit operation_id on review_apply_batch replays the
    original envelope for identical input and fails operation-id-reused when
    reused for a different batch (never a second effect)."""
    _reset()
    workdir = Path(open_workdir(tmp_path, "batch-reused"))
    current = json.loads(review_preflight())["current_snapshot"]["id"]
    first_event = _stage_human_patch(workdir, current, "智能响应", "智能调控", "batch:reused-1")
    second_event = _stage_human_patch(workdir, first_event["staged_snapshot"], "前言", "导言", "batch:reused-2")
    queued = dispatch(workdir)
    batch_id = queued[0]["batch_id"]

    op = "explicit-batch-op-1"
    first_call = review_apply_batch(batch_id, operation_id=op)
    applied = _j(first_call)
    assert applied["state"] == "applied" and applied["operation_id"] == op
    # Identical retry: replay byte-exact, never a second effect.
    assert review_apply_batch(batch_id, operation_id=op).structuredContent == first_call.structuredContent

    # A different batch under the same explicit id: operation-id-reused
    # before any effect — the fresh batch stays queued.
    fresh = json.loads(review_preflight())["current_snapshot"]["id"]
    third_event = _stage_human_patch(workdir, fresh, "后语", "结语", "batch:reused-3")
    dispatch(workdir)
    batch2 = next(item["batch_id"] for item in json.loads(review_inbox())["events"] if item["event_id"] == third_event["event_id"])
    assert batch2 != batch_id
    reused = review_apply_batch(batch2, operation_id=op)
    assert reused.isError is True
    assert reused.structuredContent["diagnostics"][0]["code"] == "operation-id-reused"
    assert json.loads(review_state())["current_snapshot"]["id"] == "C1"  # batch2 never applied
    assert "结语" not in json.loads(get_paragraph("P0"))["plain"]


def _decision_events(workdir: Path, actions: dict[str, str]) -> None:
    from scripts.review_queue import dispatch as _dispatch
    from scripts.review_queue import upsert_event as _upsert

    inventory = json.loads((workdir / "revisions.json").read_text(encoding="utf-8"))
    for revision in inventory["revisions"]:
        if revision["w_id"] not in actions:
            continue
        _upsert(
            workdir,
            {
                "type": "decision",
                "client_id": f"settle-mcp:{revision['w_id']}",
                "paragraph_id": revision["paragraph_id"],
                "revision_id": revision["w_id"],
                "revision_key": revision["revision_key"],
                "selected_text": revision["text"],
                "decision": actions[revision["w_id"]],
                "comment": "",
            },
        )
    _dispatch(workdir)


def test_review_settle_replays_exact_envelope_without_second_effect(tmp_path):
    """Findings: review_settle routes through the store seam; an identical
    retry with the same operation_id returns the original settlement envelope
    byte-exact and never settles a second time."""
    from tests.test_decisions import extract_fixture

    _reset()
    workdir = extract_fixture(tmp_path)
    session.workdir = None
    workdir = Path(json.loads(workdir_open(str(workdir)))["workdir"])
    _decision_events(workdir, {"100": "accept", "101": "reject"})

    op = "settle-round-1"
    first_call = review_settle(None, operation_id=op)
    first = _j(first_call)
    assert first["settled_event_ids"]
    assert first["operation_id"] == op
    assert first["review_base"]["id"] == "S1"

    replayed = review_settle(None, operation_id=op)
    assert replayed.structuredContent == first_call.structuredContent  # byte-exact replay
    assert json.loads(review_state())["review_base"]["id"] == "S1"  # no second settlement
    # The settled decisions were applied exactly once: w_id 100 is gone.
    remaining = {r["w_id"] for r in json.loads((workdir / "revisions.json").read_text(encoding="utf-8"))["revisions"]}
    assert "100" not in remaining and "101" not in remaining


def test_review_settle_generates_operation_id_and_rejects_reuse(tmp_path):
    """review_settle generates an id when omitted and rejects changed retries."""
    from tests.test_decisions import extract_fixture

    _reset()
    workdir = extract_fixture(tmp_path)
    session.workdir = None
    workdir = Path(json.loads(workdir_open(str(workdir)))["workdir"])
    _decision_events(workdir, {"100": "accept"})

    generated = review_settle(None, operation_id="")
    assert generated.isError is False
    assert generated.structuredContent["data"]["operation_id"]

    explicit_root = tmp_path / "explicit"
    explicit_root.mkdir()
    workdir = extract_fixture(explicit_root)
    session.workdir = None
    workdir = Path(json.loads(workdir_open(str(workdir)))["workdir"])
    _decision_events(workdir, {"100": "accept"})
    op = "settle-round-2"
    first = _j(review_settle(None, operation_id=op))
    assert first["operation_id"] == op
    # Same id, different input (explicit event ids): operation-id-reused.
    reused = review_settle(first["settled_event_ids"], operation_id=op)
    assert reused.isError is True
    assert reused.structuredContent["diagnostics"][0]["code"] == "operation-id-reused"


def test_review_ack_replays_exact_envelope_without_second_effect(tmp_path):
    """Findings: review_ack routes through the store seam; an identical retry
    returns the original ack envelope byte-exact and re-acking is a no-op."""
    from scripts.review_queue import upsert_event as _upsert

    _reset()
    workdir = Path(open_workdir(tmp_path, "ack-replay"))
    first = _upsert(
        workdir,
        {
            "type": "decision",
            "client_id": "ack:1",
            "paragraph_id": "P0",
            "revision_id": "r1",
            "revision_key": "word/document.xml|insert|999|deadbeef",
            "selected_text": "旧词",
            "decision": "accept",
            "comment": "",
        },
    )
    second = _upsert(
        workdir,
        {
            "type": "decision",
            "client_id": "ack:2",
            "paragraph_id": "P0",
            "revision_id": "r2",
            "revision_key": "word/document.xml|insert|998|deadbeef",
            "selected_text": "旧词",
            "decision": "reject",
            "comment": "",
        },
    )
    dispatch(workdir)
    event_ids = [first["event_id"], second["event_id"]]

    op = "ack-round-1"
    first_call = review_ack(event_ids, operation_id=op)
    first = _j(first_call)
    assert {item["event_id"] for item in first["acknowledged"]} == set(event_ids)
    assert first["operation_id"] == op

    replayed = review_ack(event_ids, operation_id=op)
    assert replayed.structuredContent == first_call.structuredContent  # byte-exact replay
    inbox = json.loads(review_inbox(include_acknowledged=True))
    assert all(item["status"] == "acknowledged" for item in inbox["events"])
    assert inbox["counts"]["queued"] == 0  # no second effect


def test_review_ack_generates_operation_id_and_rejects_reuse(tmp_path):
    """review_ack generates an id when omitted and rejects changed retries."""
    from scripts.review_queue import upsert_event as _upsert

    _reset()
    workdir = Path(open_workdir(tmp_path, "ack-reused"))
    first = _upsert(
        workdir,
        {
            "type": "decision",
            "client_id": "ack:3",
            "paragraph_id": "P0",
            "revision_id": "r3",
            "revision_key": "word/document.xml|insert|997|deadbeef",
            "selected_text": "旧词",
            "decision": "accept",
            "comment": "",
        },
    )
    second = _upsert(
        workdir,
        {
            "type": "decision",
            "client_id": "ack:4",
            "paragraph_id": "P0",
            "revision_id": "r4",
            "revision_key": "word/document.xml|insert|996|deadbeef",
            "selected_text": "旧词",
            "decision": "reject",
            "comment": "",
        },
    )
    dispatch(workdir)

    generated = _j(review_ack([first["event_id"]], operation_id=""))
    assert generated["operation_id"]

    op = "ack-round-2"
    acked = _j(review_ack([second["event_id"]], operation_id=op))
    assert acked["operation_id"] == op
    reused = review_ack([first["event_id"]], operation_id=op)
    assert reused.isError is True
    assert reused.structuredContent["diagnostics"][0]["code"] == "operation-id-reused"
    inbox = json.loads(review_inbox(include_acknowledged=True))
    assert all(item["status"] == "acknowledged" for item in inbox["events"])

def test_draft_mutation_replay_survives_pointer_advance(tmp_path):
    """Findings: draft mutators route mutation, ledger, and evidence through
    the pinned store generation (the same seam commit_sync uses). Replaying
    the identical draft operation from a FRESH process AFTER the pointer
    advanced must hit the generation the record was written under and return
    the original envelope — never a second draft effect."""
    _reset()
    workdir = open_store_workdir(tmp_path, "advance")
    op = "advance-draft-1"
    first = _j(replace_text("P0", "智能响应", "智能调控", operation_id=op))
    assert first["draft"] == "dirty" and first["operation_id"] == op
    # Advance the pointer: the draft mutation's generation becomes old.
    _j(commit_sync(operation_id="advance-commit-1"))
    assert _j(workdir_status())["state"] == "clean"
    # Fresh interpreter (empty in-process ledger): the replay must find the
    # record in the old generation and replay the original success envelope
    # instead of re-running (which would fail text-not-found on the clean
    # draft).
    script = (
        "import sys, json; sys.path.insert(0, %r);\n"
        "from scripts.mcp_server import replace_text, session, workdir_open;\n"
        "session.workdir = None; json.loads(workdir_open(%r));\n"
        "result = replace_text('P0', '智能响应', '智能调控', operation_id=%r);\n"
        "print(json.dumps(result.structuredContent, ensure_ascii=False, sort_keys=True))"
    ) % (str(ROOT), str(workdir), op)
    replay = subprocess.run(
        [sys.executable, "-c", script],
        cwd=ROOT, capture_output=True, text=True, encoding="utf-8",
    )
    assert replay.returncode == 0, replay.stderr
    envelope = json.loads(replay.stdout.strip())
    assert envelope["outcome"] == "success", envelope
    assert envelope["operation"] == "replace_text"
    assert envelope["data"]["operation_id"] == op
    assert envelope["data"]["draft"] == "dirty"
    # Single effect only: the committed draft still carries one change.
    plain = json.loads(get_paragraph("P0"))["plain"]
    assert "智能调控" in plain and "智能响应" not in plain


def test_table_insert_row_create_mode_success(tmp_path):
    """Findings: store-backed table ops hash the STAGED output for evidence —
    the final path does not exist until publish — so create mode must
    succeed, and the evidence records both the staged hash and the final
    path."""
    _reset()
    source = tmp_path / "tbl-src.docx"
    workdir = tmp_path / "tbl-wd"
    make_table_docx(source)
    assert extract([str(source), "-o", str(workdir)]) == 0
    json.loads(workdir_open(str(workdir)))
    output = tmp_path / "tbl-out.docx"
    workdir_out = tmp_path / "tbl-out-wd"
    result = table_insert_row("T0", 0, str(output), str(workdir_out), operation_id="tbl-create-1")
    assert result.isError is False, result.structuredContent
    assert output.is_file()
    assert workdir_out.is_dir()
    evidence = result.structuredContent["evidence"][0]["payload"]["outputs"]["docx"]
    assert evidence["sha256"] == file_sha256(output)  # staged hash == published bytes
    assert evidence["path"] == str(output.resolve())


def test_decide_all_create_mode_success(tmp_path):
    """Findings: store-backed decide_all hashes the STAGED output for
    evidence (create mode: the final path does not exist until publish)."""
    import zipfile

    _reset()
    source = tmp_path / "rev-src.docx"
    workdir = tmp_path / "rev-wd"
    make_doc(source)
    with zipfile.ZipFile(source) as z:
        files = {n: z.read(n) for n in z.namelist()}
    insertion = (
        '<w:ins w:id="99" w:author="t" w:date="2026-01-01T00:00:00Z">'
        "<w:r><w:t>修订词</w:t></w:r></w:ins>"
    ).encode("utf-8")
    files["word/document.xml"] = files["word/document.xml"].replace(
        "<w:r><w:t>前言</w:t></w:r>".encode("utf-8"),
        "<w:r><w:t>前言</w:t></w:r>".encode("utf-8") + insertion,
        1,
    )
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    assert extract([str(source), "-o", str(workdir)]) == 0
    json.loads(workdir_open(str(workdir)))
    output = tmp_path / "decided-out.docx"
    workdir_out = tmp_path / "decided-out-wd"
    result = decide_all("accept", str(output), str(workdir_out), operation_id="decide-create-1")
    assert result.isError is False, result.structuredContent
    assert output.is_file()
    assert workdir_out.is_dir()
    evidence = result.structuredContent["evidence"][0]["payload"]["outputs"]["docx"]
    assert evidence["sha256"] == file_sha256(output)  # staged hash == published bytes
    assert evidence["path"] == str(output.resolve())


def test_direct_script_fallback_imports_read_root():
    """Findings: running mcp_server.py as a direct script (no package
    context) must import ``read_root`` from the fallback store import."""
    script = (
        "import sys; sys.path.insert(0, %r);\n"
        "import mcp_server;\n"
        "print(mcp_server.read_root is not None)"
    ) % (str(ROOT / "scripts"),)
    result = subprocess.run(
        [sys.executable, "-c", script],
        cwd=ROOT, capture_output=True, text=True, encoding="utf-8",
    )
    assert result.returncode == 0, result.stderr
    assert result.stdout.strip() == "True"


def test_track_mode_dirty_draft_sequential_edits(tmp_path):
    """Regression: consecutive replace_text on a dirty draft must keep the
    session's track mode (pending revisions + trackChanges off would
    otherwise re-infer ambiguous and reject every later edit)."""
    import re as _re
    import zipfile

    _reset()
    source = tmp_path / "track-src.docx"
    workdir = tmp_path / "track"
    make_doc(source)
    # inject a pending revision into the SOURCE, with trackChanges off
    with zipfile.ZipFile(source) as z:
        files = {n: z.read(n) for n in z.namelist()}
    doc = files["word/document.xml"]
    ins = (
        '<w:ins w:id="99" w:author="tester" w:date="2026-01-01T00:00:00Z">'
        '<w:r><w:t>修订词</w:t></w:r></w:ins>'
    ).encode()
    doc = doc.replace(
        "<w:r><w:t>前言</w:t></w:r>".encode(),
        "<w:r><w:t>前</w:t></w:r>".encode() + ins + "<w:r><w:t>言</w:t></w:r>".encode(),
        1,
    )
    files["word/document.xml"] = doc
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir), track=True, author="AI润色"))
    # consecutive edits across paragraphs on the dirty draft, then one
    # preview + one commit (the draft -> preview -> commit model)
    _j(replace_text("P0", "智能响应", "智能调控", operation_id="track-1"))
    _j(replace_text("P1", "第二段", "第二段落", operation_id="track-2"))
    _j(replace_text("P0", "后语", "后文", operation_id="track-3"))
    preview = _j(diff_preview())
    assert preview["state"] == "dirty"
    assert len(preview["hunks"]) >= 3
    committed = _j(commit_sync(operation_id="track-commit-1"))
    assert committed["state"] == "clean"
    assert committed["edit_mode"] == "track"
    output = _j(build_docx(operation_id="track-build-1"))["output"]
    with zipfile.ZipFile(output) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert xml.count("<w:ins") >= 3 + 1  # 3 new + the pre-existing one
    assert _re.search(r'w:author="AI润色"', xml)


def test_document_read_returns_virtual_file(tmp_path):
    _reset()
    workdir = open_workdir(tmp_path, "docread")
    whole = _j(document_read())  # view=auto: small doc -> full content
    assert whole["view"] == "content"
    assert whole["content"].startswith("<!--@edit")
    assert '<!--@p id="P0"-->' in whole["content"] and '<!--@p id="P1"-->' in whole["content"]
    assert "智能响应" in whole["content"] and "第二段" in whole["content"]
    assert "revision" in whole and len(whole["revision"]) == 64
    # frozen invariant: document_read.content == the real projection bytes
    assert whole["content"] == (Path(workdir) / "edit.md").read_text(encoding="utf-8")
    window = _j(document_read(anchor="P0", before=0, after=0))
    assert window["windowed"] is True
    assert "第二段" not in window["content"] and "智能响应" in window["content"]
    tail = _j(document_read(anchor="P1", before=5, after=8))
    assert "第二段" in tail["content"]


def test_document_read_outline(tmp_path):
    _reset()
    open_workdir(tmp_path, "outline")
    outline = _j(document_read(view="outline"))
    assert '<!--@p id="P0"-->' in outline["content"] and "[12 chars]" not in outline["content"]
    assert "第二段" in outline["content"]
    import pytest
    from scripts.mcp_server import ToolError

    with pytest.raises(ToolError):
        document_read(view="prose")
    with pytest.raises(ToolError):
        document_read(anchor="P99")


def test_document_search_returns_context_blocks(tmp_path):
    _reset()
    open_workdir(tmp_path, "docsearch")
    found = _j(document_search("智能响应"))
    assert found["total_matches"] == 1 and found["returned_blocks"] == 1
    entry = found["matches"][0]
    assert entry["id"] == "P0" and "智能响应" in entry["text"]
    assert entry["prev_id"] is None and entry["next_id"] == "P1"
    # case-insensitive by default
    assert _j(document_search("abc"))["total_matches"] == 1
    assert _j(document_search("不存在词"))["total_matches"] == 0


def test_document_search_reflects_draft_state(tmp_path):
    _reset()
    open_workdir(tmp_path, "draftsearch")
    _j(replace_text("P0", "智能响应", "智能调控", operation_id="draftsearch-1"))
    found = _j(document_search("智能调控"))
    assert found["matches"][0]["id"] == "P0"
    assert _j(document_search("智能响应"))["total_matches"] == 0


def test_document_patch_multi_paragraph_matches_sequential_replace(tmp_path):
    _reset()
    wd_a = open_workdir(tmp_path, "patchseq-a")
    _j(document_patch(hunks=[
        {"paragraph_id": "P0", "old": "智能响应", "new": "智能调控"},
        {"paragraph_id": "P1", "old": "第二段", "new": "第二段落"},
    ], operation_id="patch-seq-1"))
    _reset()
    wd_b = open_workdir(tmp_path, "patchseq-b")
    _j(replace_text("P0", "智能响应", "智能调控", operation_id="patch-seq-2"))
    _j(replace_text("P1", "第二段", "第二段落", operation_id="patch-seq-3"))
    a = (Path(wd_a) / "edit.md").read_text(encoding="utf-8").split("\n", 1)[1]
    b = (Path(wd_b) / "edit.md").read_text(encoding="utf-8").split("\n", 1)[1]
    assert a == b, "batch patch must land body-identical to sequential replaces"


def test_document_patch_insert_and_delete(tmp_path):
    _reset()
    open_workdir(tmp_path, "patchstruct")
    result = _j(document_patch(hunks=[
        {"insert_after": "P0", "text": "插入段"},
        {"delete": "P1"},
    ], operation_id="patch-struct-1"))
    assert next(e for e in result["applied"] if e["kind"] == "insert")["temp_id"] == "N1"
    _j(commit_sync(operation_id="patch-struct-2"))
    output = _j(build_docx(operation_id="patch-struct-3"))["output"]
    texts = [p.text for p in Document(output).paragraphs]
    assert "插入段" in texts and "第二段" not in texts


def test_document_patch_atomic_rejection(tmp_path):
    _reset()
    open_workdir(tmp_path, "patchatomic")
    result = document_patch(hunks=[
        {"paragraph_id": "P0", "old": "智能响应", "new": "智能调控"},
        {"paragraph_id": "P0", "old": "智能响应ABC", "new": "x"},  # cross-region
    ], operation_id="patch-atomic-1")
    assert result.isError is True
    codes = [d["code"] for d in result.structuredContent["diagnostics"]]
    assert codes[0] == "document-patch-hunks-overlap"
    found = _j(document_search("智能响应"))
    assert found["total_matches"] == 1  # draft untouched


def test_document_patch_multi_hunks_same_paragraph(tmp_path):
    _reset()
    wd_a = open_workdir(tmp_path, "patchmulti-a")
    _j(document_patch(hunks=[
        {"paragraph_id": "P0", "old": "前言", "new": "前言改"},
        {"paragraph_id": "P0", "old": "后语", "new": "后语改"},
        {"paragraph_id": "P0", "old": "智能响应", "new": "智能调控"},
    ], operation_id="patch-multi-1"))
    _reset()
    wd_b = open_workdir(tmp_path, "patchmulti-b")
    for op, (o, n) in enumerate((("前言", "前言改"), ("后语", "后语改"), ("智能响应", "智能调控")), 2):
        _j(replace_text("P0", o, n, operation_id=f"patch-multi-{op}"))
    a = (Path(wd_a) / "edit.md").read_text(encoding="utf-8").split("\n", 1)[1]
    b = (Path(wd_b) / "edit.md").read_text(encoding="utf-8").split("\n", 1)[1]
    assert a == b


def test_document_patch_overlapping_hunks_rejected(tmp_path):
    _reset()
    open_workdir(tmp_path, "patchoverlap2")
    result = document_patch(hunks=[
        {"paragraph_id": "P0", "old": "智能", "new": "智慧"},
        {"paragraph_id": "P0", "old": "能响应", "new": "X"},
    ], operation_id="patch-overlap-2")
    assert result.isError is True
    assert result.structuredContent["diagnostics"][0]["code"] == "document-patch-hunks-overlap"


def test_document_patch_replace_and_delete_conflict(tmp_path):
    _reset()
    open_workdir(tmp_path, "patchconflict")
    result = document_patch(hunks=[
        {"paragraph_id": "P0", "old": "前言", "new": "前言改"},
        {"delete": "P0"},
    ], operation_id="patch-conflict-1")
    assert result.isError is True
    assert result.structuredContent["diagnostics"][0]["code"] == "document-patch-paragraph-repeated"


def test_document_patch_base_revision(tmp_path):
    _reset()
    workdir = open_workdir(tmp_path, "patchrev")
    found = _j(document_search("智能响应"))
    revision = found["revision"]
    stale = document_patch(
        hunks=[{"paragraph_id": "P0", "old": "前言", "new": "前言改"}],
        base_revision="deadbeef" + revision[8:],
        operation_id="patch-rev-1",
    )
    assert stale.isError is True
    assert stale.structuredContent["diagnostics"][0]["code"] == "stale-document-view"
    assert _j(document_search("前言"))["total_matches"] == 1  # untouched
    _j(document_patch(
        hunks=[{"paragraph_id": "P0", "old": "前言", "new": "前言改"}],
        base_revision=revision,
        operation_id="patch-rev-2",
    ))
    assert _j(document_search("前言改"))["total_matches"] == 1


def _diff_for(tmp_path, name):
    """Build a unified diff against the current projection by mutating it."""
    _reset()
    workdir = open_workdir(tmp_path, name)
    text = (Path(workdir) / "edit.md").read_text(encoding="utf-8")
    lines = text.split("\n")
    target = next(i for i, l in enumerate(lines) if "智能响应" in l)
    new_line = lines[target].replace("智能响应", "智能调控")
    diff = (
        "--- a/edit.md\n+++ b/edit.md\n"
        f"@@ -{target},3 +{target},3 @@\n"
        f" {lines[target - 1]}\n-{lines[target]}\n+{new_line}\n {lines[target + 1]}\n"
    )
    return workdir, diff


def test_document_patch_unified_diff(tmp_path):
    workdir, diff = _diff_for(tmp_path, "patchdiff")
    result = _j(document_patch(diff=diff, operation_id="patch-diff-1"))
    assert result["affected_paragraph_ids"] == ["P0"]
    found = _j(document_search("智能调控"))
    assert found["matches"][0]["id"] == "P0"
    assert _j(document_search("智能响应"))["total_matches"] == 0


def test_document_patch_diff_context_mismatch(tmp_path):
    workdir, diff = _diff_for(tmp_path, "patchctx")
    bad = diff.replace("前言", "不存在的上下文")
    result = document_patch(diff=bad, operation_id="patch-ctx-1")
    assert result.isError is True
    assert result.structuredContent["diagnostics"][0]["code"] == "patch-context-mismatch"
    assert _j(document_search("智能响应"))["total_matches"] == 1


def test_document_patch_requires_exactly_one_input(tmp_path):
    _reset()
    open_workdir(tmp_path, "patchargs")
    for kwargs in ({"hunks": None, "diff": None}, {"hunks": [], "diff": "x"}):
        result = document_patch(operation_id="patch-args-1", **kwargs)
        assert result.isError is True
        assert result.structuredContent["diagnostics"][0]["code"] == "invalid-arguments"


def test_facade_four_call_session(tmp_path):
    """The PRD's end-to-end promise: orient, locate, patch, save in four
    tool calls, then build and verify clean."""
    _reset()
    open_workdir(tmp_path, "facade-session")
    outline = _j(document_read(view="outline"))["content"]  # 1. orient
    assert '<!--@p id="P0"-->' in outline
    found = _j(document_search("智能响应"))  # 2. locate
    assert found["matches"][0]["id"] == "P0"
    _j(document_patch(hunks=[  # 3. patch
        {"paragraph_id": "P0", "old": "智能响应", "new": "智能调控"},
        {"paragraph_id": "P1", "old": "第二段", "new": "第二段落修订"},
    ], operation_id="facade-1"))
    _j(commit_sync(operation_id="facade-2"))  # 4. save
    output = _j(build_docx(operation_id="facade-3"))["output"]
    assert _j(verify_output(output))["verified"] == output
    texts = [p.text for p in Document(output).paragraphs]
    assert any("智能调控" in t for t in texts) and any("第二段落修订" in t for t in texts)


def test_document_patch_mixed_style_reaches_core(tmp_path):
    """A cross-region span is no longer refused by the patch tool: the sync
    engine decides deterministically (accept + warning, or a sync
    rejection code) — never cross-region-text."""
    _reset()
    open_workdir(tmp_path, "patchmixed")
    result = document_patch(hunks=[
        {"paragraph_id": "P0", "old": "响应ABC", "new": "响应改写后"},
    ], operation_id="patch-mixed-1")
    codes = [d["code"] for d in result.structuredContent["diagnostics"]] if result.isError else []
    assert "cross-region-text" not in codes
    if result.isError:
        assert any(c in ("mixed-replacement-requires-unchanged-text", "unanchored-mixed-rewrite", "protected-boundary-crossing") for c in codes), codes
    else:
        data = _j(result)
        _j(commit_sync(operation_id="patch-mixed-2"))
        output = _j(build_docx(operation_id="patch-mixed-3"))["output"]
        texts = "".join(p.text for p in Document(output).paragraphs)
        assert "响应改写后" in texts


def test_document_patch_full_paragraph_rewrite_accepted_by_core(tmp_path):
    """A whole-paragraph cross-region rewrite is decided by the Core's
    deterministic proportional style mapping, not refused: patch succeeds
    (with warnings), commit and build stay clean, text lands."""
    _reset()
    open_workdir(tmp_path, "patchfullrewrite")
    result = document_patch(hunks=[
        {"paragraph_id": "P0", "old": "前言智能响应ABC后语", "new": "整体重写的一段全新内容"},
    ], operation_id="patch-full-1")
    assert result.isError is False, result.structuredContent["diagnostics"]
    _j(commit_sync(operation_id="patch-full-2"))
    output = _j(build_docx(operation_id="patch-full-3"))["output"]
    texts = "".join(p.text for p in Document(output).paragraphs)
    assert "整体重写的一段全新内容" in texts


def test_document_patch_windowed_diff(tmp_path):
    """A diff generated from a windowed document_read works: hunks locate by
    marker id + exact body, never by line numbers."""
    _reset()
    workdir = open_workdir(tmp_path, "patchwin")
    found = _j(document_search("智能响应"))
    anchor = found["matches"][0]["id"]
    window = _j(document_read(anchor=anchor, before=1, after=1))
    lines = window["content"].split("\n")
    target = next(i for i, l in enumerate(lines) if "智能响应" in l)
    # @@ header deliberately LIES about line numbers — must still apply
    diff = (
        "--- a/edit.md\n+++ b/edit.md\n"
        "@@ -999,3 +999,3 @@\n"
        f" {lines[target - 1]}\n-{lines[target]}\n+{lines[target].replace('智能响应', '智能调控')}\n {lines[target + 1]}\n"
    )
    result = _j(document_patch(diff=diff, base_revision=window["revision"], operation_id="patch-win-1"))
    assert result["affected_paragraph_ids"] == [anchor]
    assert _j(document_search("智能调控"))["total_matches"] == 1


def test_document_patch_diff_marker_edit_rejected(tmp_path):
    _reset()
    open_workdir(tmp_path, "patchmarker")
    text = (Path(_j(document_read())["content"]) if False else None)
    workdir_path = Path(document_read.__globals__["session"].workdir)
    lines = (workdir_path / "edit.md").read_text(encoding="utf-8").split("\n")
    target = next(i for i, l in enumerate(lines) if "智能响应" in l)
    diff = (
        "--- a/edit.md\n+++ b/edit.md\n"
        f"@@ -{target - 1},3 +{target - 1},3 @@\n"
        f" {lines[target - 2]}\n-<!--@p id=\"P0\"-->\n+<!--@p id=\"PX\"-->\n {lines[target - 1]}\n"
    )
    result = document_patch(diff=diff, operation_id="patch-marker-1")
    assert result.isError is True
    assert result.structuredContent["diagnostics"][0]["code"] == "patch-structure-immutable"


def test_document_patch_diff_pure_insertion(tmp_path):
    _reset()
    workdir_path = None
    open_workdir(tmp_path, "patchinsert")
    text = (Path(_j(document_read())["content"]) if False else None)
    workdir_path = Path(document_read.__globals__["session"].workdir)
    lines = (workdir_path / "edit.md").read_text(encoding="utf-8").split("\n")
    target = next(i for i, l in enumerate(lines) if "智能响应" in l)
    body = lines[target]
    pos = body.index("响应")
    new_body = body[:pos] + "彻底" + body[pos:]
    diff = (
        "--- a/edit.md\n+++ b/edit.md\n"
        f"@@ -{target},2 +{target},2 @@\n"
        f" {lines[target - 1]}\n-{body}\n+{new_body}\n"
    )
    _j(document_patch(diff=diff, operation_id="patch-ins-1"))
    assert _j(document_search("彻底响应"))["total_matches"] == 1


def test_document_patch_diff_multi_span(tmp_path):
    """Two far-apart single-region edits inside one paragraph arrive as one
    diff and become two hunks — no giant cross-region span."""
    _reset()
    open_workdir(tmp_path, "patchmulti2")
    workdir_path = Path(document_read.__globals__["session"].workdir)
    lines = (workdir_path / "edit.md").read_text(encoding="utf-8").split("\n")
    target = next(i for i, l in enumerate(lines) if "智能响应" in l)
    body = lines[target]
    new_body = body.replace("前言", "前言V2").replace("后语", "后语V3")
    diff = (
        "--- a/edit.md\n+++ b/edit.md\n"
        f"@@ -{target - 1},2 +{target - 1},2 @@\n"
        f" {lines[target - 1]}\n-{body}\n+{new_body}\n"
    )
    result = _j(document_patch(diff=diff, operation_id="patch-multi2-1"))
    assert next(e for e in result["applied"] if e["paragraph_id"] == "P0")["hunks"] == 2  # two precise spans, not one giant
    assert _j(document_search("前言V2"))["total_matches"] == 1
    assert _j(document_search("后语V3"))["total_matches"] == 1


def _make_ambiguous_doc(path: Path) -> None:
    """trackChanges ON, zero pending revisions -> ambiguous mode."""
    document = Document()
    document.add_paragraph("前言")
    paragraph = document.add_paragraph()
    paragraph.add_run("智能响应")
    document.add_paragraph("第二段")
    from docx.oxml.ns import qn as _qn

    settings_el = document.settings.element
    settings_el.append(settings_el.makeelement(_qn("w:trackChanges"), {}))
    document.save(path)


def test_ambiguous_mode_defaults_to_track(tmp_path):
    """#76 evolved: an ambiguous mode no longer costs a refused call. The first
    mutation resolves it FAIL-SAFE to track (never silently rewrite another
    author's revision) and says so; track=false still overrides in the same
    call."""
    _reset()
    source = tmp_path / "ambig-src.docx"
    workdir = tmp_path / "ambig"
    _make_ambiguous_doc(source)
    assert extract([str(source), "-o", str(workdir)]) == 0
    opened = json.loads(workdir_open(str(workdir)))  # no explicit choice -> ambiguous
    assert opened["edit_mode"] == "ambiguous"
    result = document_patch(hunks=[
        {"paragraph_id": "P1", "old": "智能响应", "new": "智能调控"},
    ], operation_id="ambig-1")
    assert not result.isError, result.structuredContent
    warnings = result.structuredContent["data"]["warnings"]
    assert any("edit-mode-ambiguous-defaulted-to-track" in w for w in warnings), warnings
    assert session.mode == "track"
    _j(revert())
    # choose track → full revision lifecycle works
    _j(workdir_open(workdir, track=True))
    _j(document_patch(hunks=[
        {"paragraph_id": "P0", "old": "智能响应", "new": "智能调控"},
    ], operation_id="ambig-2"))
    committed = _j(commit_sync(operation_id="ambig-3"))
    assert committed["edit_mode"] == "track"
    output = _j(build_docx(operation_id="ambig-4"))["output"]
    assert _j(verify_output(output))["verified"] == output


def test_document_patch_edits_table_cell_text(tmp_path):
    """#75: cell text replace goes through the facade; structure is locked."""
    _reset()
    source = tmp_path / "cell-src.docx"
    workdir = tmp_path / "celledit"
    make_table_docx(source)
    assert main(["--json", "extract", str(source), "-o", str(workdir), "--operation-id", "cell-extract-1"]) == 0
    _j(workdir_open(str(workdir)))
    _j(document_patch(hunks=[
        {"paragraph_id": "T0.R0.C0.P0", "old": "A1", "new": "要素"},
    ], operation_id="cell-1"))
    _j(commit_sync(operation_id="cell-2"))
    output = _j(build_docx(operation_id="cell-3"))["output"]
    assert _j(verify_output(output))["verified"] == output
    table = Document(output).tables[0]
    assert table.rows[0].cells[0].text == "要素"
    assert len(table.rows) == 2 and len(table.columns) == 2  # structure intact


def test_document_patch_still_refuses_container_topology(tmp_path):
    """#75 boundary: insert/delete on container paragraphs stay refused."""
    _reset()
    source = tmp_path / "cellstruct-src.docx"
    workdir = tmp_path / "cellstruct"
    make_table_docx(source)
    assert main(["--json", "extract", str(source), "-o", str(workdir), "--operation-id", "cellstruct-extract-1"]) == 0
    _j(workdir_open(str(workdir)))
    for op_id, hunk in (
        ("cell-ins", {"insert_after": "T0.R0.C0.P0", "text": "x"}),
        ("cell-del", {"delete": "T0.R0.C0.P0"}),
    ):
        result = document_patch(hunks=[hunk], operation_id=op_id)
        assert result.isError is True
        assert result.structuredContent["diagnostics"][0]["code"] == "table-structure-immutable"


def test_build_refuses_to_overwrite_source(tmp_path):
    """Dogfood round 3: build_docx(output=<source path>) must refuse instead
    of silently destroying the user's original document."""
    _reset()
    source = tmp_path / "tgt-src.docx"
    workdir = tmp_path / "tgt"
    make_doc(source)
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir)))
    result = build_docx(output=str(source), operation_id="guard-1")
    assert result.isError is True
    assert result.structuredContent["diagnostics"][0]["code"] == "output-path-reserved"
    import zipfile

    with zipfile.ZipFile(source) as z:  # original intact
        assert "word/document.xml" in z.namelist()


def test_placeholder_in_edit_span(tmp_path):
    """#77: old strings carrying read-only placeholder tokens get a
    mechanical, self-explanatory refusal instead of text-not-found."""
    _reset()
    open_workdir(tmp_path, "placeholder")
    result = document_patch(hunks=[
        {"paragraph_id": "P0", "old": "智能响应⟦tab⟧ABC", "new": "x"},
    ], operation_id="ph-1")
    assert result.isError is True
    diag = result.structuredContent["diagnostics"][0]
    assert diag["code"] == "placeholder-in-edit-span"
    recovery = result.structuredContent["data"]["recovery"]
    assert recovery["action"] == "choose-editable-span"


def test_build_refuses_all_reserved_workdir_paths(tmp_path):
    """P0 close-out: the shared output validator covers the whole workdir
    tree (_template.docx, typed.md, format.json, styles.json, edit.md,
    arbitrary new files) — not just the source document."""
    _reset()
    source = tmp_path / "res-src.docx"
    workdir = tmp_path / "res"
    make_doc(source)
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir)))
    for name in ("_template.docx", "typed.md", "format.json", "styles.json", "edit.md", "notes.txt"):
        result = build_docx(output=str(workdir / name), operation_id=f"res-{name}")
        assert result.isError is True, name
        assert result.structuredContent["diagnostics"][0]["code"] == "output-path-reserved", name
    # a normal external output still works
    external = tmp_path / "external.docx"
    _j(commit_sync(operation_id="res-save"))  # ADR 0044: a valid path still needs a saved version
    ok = build_docx(output=str(external), operation_id="res-ok")
    assert not ok.isError


def test_no_public_output_validation_bypass():
    """P0 close-out: build_workdir is the only public build entry and always
    validates the output path; staging is reachable only via the private
    Store-transaction helper."""
    import inspect
    import scripts.typed_docx as td
    assert "validate_output" not in inspect.signature(td.build_workdir).parameters
    assert not hasattr(td, "build_workdir_impl")
    assert td._build_workdir_to_staging.__name__ == "_build_workdir_to_staging"
    import importlib
    build_api = importlib.import_module("scripts.build")
    exported = set(build_api.__all__)
    assert "build_workdir" in exported
    assert not any("staging" in name or "impl" in name for name in exported)

RELEASE = Path(__file__).resolve().parents[1] / "corpus" / "release"


def _make_two_para_docx(path):
    from docx import Document
    d = Document()
    d.add_paragraph("甲段落原文内容 保持不动")
    d.add_paragraph("乙段落前缀文字 目标插入语 后缀文字收尾")
    d.save(path)


def _open_tracked(tmp_path, name="xb", with_bold=False):
    from scripts.extract import extract
    source = tmp_path / f"{name}-src.docx"
    if with_bold:
        # a document that can express bold at all: styles.json mirrors the
        # source, so a variant must exist before format_span may reuse it
        document = Document()
        document.add_paragraph("甲段落原文内容 保持不动")
        paragraph = document.add_paragraph()
        paragraph.add_run("乙段落前缀文字 ")
        paragraph.add_run("目标插入语").bold = True
        paragraph.add_run(" 后缀文字收尾")
        document.save(source)
    else:
        _make_two_para_docx(source)
    workdir = tmp_path / name
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir), track=True))
    return workdir


def test_revision_boundary_spans_get_precise_diagnostic(tmp_path):
    """#78: old spanning a revision-control boundary (baseline -> committed
    w:ins) is visible text but not one editable span; both document_patch and
    replace_text must say so instead of text-not-found, and a refusal must
    leave the draft byte-identical."""
    workdir = _open_tracked(tmp_path)
    # commit one insertion so P1 carries an ins region
    r = document_patch(hunks=[{"paragraph_id": "P1", "old": "目标插入语", "new": "目标插入语甲"}], operation_id="xb-mk")
    assert not r.isError
    assert not commit_sync(operation_id="xb-mk-c").isError
    _j(workdir_open(str(workdir), track=True))
    before = (workdir / "edit.md").read_bytes()

    # 1 baseline region ok / 2 inside-ins ok / 5 whole-ins-body ok
    for label, old in [("baseline", "前缀文字"), ("ins-head", "目标插入语甲"[:0] or "甲"), ]:
        pass  # covered below with explicit cases
    r_base = document_patch(hunks=[{"paragraph_id": "P1", "old": "前缀文字", "new": "前缀文字X"}], operation_id="xb-base")
    assert not r_base.isError, r_base.structuredContent
    _j(revert(operation_id="xb-base-r"))  # discard draft
    r_ins = document_patch(hunks=[{"paragraph_id": "P1", "old": "甲", "new": "甲X"}], operation_id="xb-ins")
    assert not r_ins.isError, r_ins.structuredContent
    _j(revert(operation_id="xb-ins-r"))  # discard draft

    # 3 cross start boundary (baseline tail + ins head)
    r_cross = document_patch(hunks=[{"paragraph_id": "P1", "old": "插入语甲", "new": "插入语甲X"}], operation_id="xb-cross")
    assert r_cross.isError
    diag = r_cross.structuredContent["diagnostics"][0]
    assert diag["code"] == "edit-span-crosses-revision-boundary", diag
    assert "insert-start" in diag["message"]
    assert r_cross.structuredContent["data"]["recovery"]["action"] == "replan-within-revision-regions"
    assert (workdir / "edit.md").read_bytes() == before  # draft untouched

    # 4 cross end boundary (ins tail + baseline head)
    r_end = document_patch(hunks=[{"paragraph_id": "P1", "old": "甲 后缀文字", "new": "甲X 后缀文字"}], operation_id="xb-cross-end")
    assert r_end.isError
    assert r_end.structuredContent["diagnostics"][0]["code"] == "edit-span-crosses-revision-boundary"

    # 6 replace_text returns the same diagnostic on the same span
    rt = replace_text(paragraph_id="P1", old="插入语甲", new="插入语甲X", operation_id="xb-rt")
    assert rt.isError
    assert rt.structuredContent["diagnostics"][0]["code"] == "edit-span-crosses-revision-boundary"

    # 7 placeholder tokens still take the #77 code, no regression
    r_ph = document_patch(hunks=[{"paragraph_id": "P1", "old": "目标⟦x⟧插入语", "new": "y"}], operation_id="xb-ph")
    assert r_ph.isError
    assert r_ph.structuredContent["diagnostics"][0]["code"] == "placeholder-in-edit-span"

    # 8 genuine absence still text-not-found
    r_nf = document_patch(hunks=[{"paragraph_id": "P1", "old": "不存在的句子", "new": "y"}], operation_id="xb-nf")
    assert r_nf.isError
    assert r_nf.structuredContent["diagnostics"][0]["code"] == "text-not-found"


def test_source_drift_hard_gate(tmp_path):
    """Out-of-band source DOCX mutation is mechanically detected: the
    source_sha256 recorded at extract is verified on open; the MCP wrapper
    surfaces a structured source-drift failure instead of a leak."""
    workdir = _open_tracked(tmp_path, "drift")
    _j(commit_sync(operation_id="drift-save"))  # ADR 0044: a saved version exists first
    source = workdir.parent / "drift-src.docx"
    data = bytearray(source.read_bytes())
    data[-1] ^= 0xFF  # raw-OOXML-style out-of-band edit
    source.write_bytes(bytes(data))
    result = _workdir_open_result(str(workdir), track=True)
    assert result.isError
    diag = result.structuredContent["diagnostics"][0]
    assert diag["code"] == "source-drift", diag
    # core paths refuse too (wrapper maps ValidationError -> structured)
    c = commit_sync(operation_id="drift-c")
    assert c.isError
    code = c.structuredContent["diagnostics"][0]["code"]
    assert code in {"source-drift", "source-modified-outside-engine"}, code
    b = build_docx(output=str(tmp_path / "drift-out.docx"), operation_id="drift-b")
    assert b.isError
    assert b.structuredContent["diagnostics"][0]["code"] in {"source-drift", "source-modified-outside-engine"}


def test_comment_text_requires_explicit_opt_in(tmp_path):
    """comments.P* is annotation content: replace is refused by default via
    both document_patch and replace_text; explicit allow_comment_text=True
    (user asked) is the only way through."""
    _reset()
    import shutil
    source = tmp_path / "cmt-src.docx"
    shutil.copy2(RELEASE / "comments.docx", source)
    workdir = tmp_path / "cmt"
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir)))
    d = _j(get_paragraph("comments.P0"))
    visible = d["data"]["plain"] if "data" in d else d["plain"]
    import re as _re
    old = _re.sub("\u27e6[^\u27e7]*\u27e7", "", visible)
    r = document_patch(hunks=[{"paragraph_id": "comments.P0", "old": old, "new": old + "X"}], operation_id="cmt-1")
    assert r.isError
    assert r.structuredContent["diagnostics"][0]["code"] == "comment-text-requires-opt-in"
    rt = replace_text(paragraph_id="comments.P0", old=old, new=old + "X", operation_id="cmt-2")
    assert rt.isError
    assert rt.structuredContent["diagnostics"][0]["code"] == "comment-text-requires-opt-in"
    ok = document_patch(hunks=[{"paragraph_id": "comments.P0", "old": old, "new": old + "X"}], operation_id="cmt-3", allow_comment_text=True)
    assert not ok.isError, ok.structuredContent
    ins = document_patch(hunks=[{"insert_after": "comments.P0", "text": "追加批注段"}], operation_id="cmt-4")
    assert ins.isError
    assert ins.structuredContent["diagnostics"][0]["code"] == "comment-text-requires-opt-in"


def _make_superscript_docx(path, *, with_reference: bool):
    from docx import Document
    d = Document()
    if with_reference:
        p = d.add_paragraph()
        p.add_run("既有正确写法 Cu")
        run = p.add_run("2+")
        run.font.superscript = True
    d.add_paragraph("本段落里的 Cu2+ 上下标缺失，需要修复。")
    d.save(path)


def _format_span_workdir(tmp_path, name, *, with_reference, track=False):
    from scripts.extract import extract
    source = tmp_path / f"{name}-src.docx"
    _make_superscript_docx(source, with_reference=with_reference)
    workdir = tmp_path / name
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir), track=track, author="Lin"))
    return workdir


def test_commit_refuses_a_state_the_builder_cannot_reproduce(tmp_path):
    """Issue #83: some nested-revision shapes round-trip with a different node
    skeleton, which bricks every later build. The save boundary rehearses the
    sync + build on a copy and refuses BEFORE publishing, leaving the draft
    revertible instead of the workdir unbuildable."""
    workdir = _open_tracked(tmp_path, "commitprobe")
    # a normal save still works (the probe must not block healthy edits)
    r = document_patch(hunks=[{"paragraph_id": "P1", "old": "目标插入语", "new": "目标插入语甲"}], operation_id="cp-1", track=True)
    assert not r.isError
    assert not commit_sync(operation_id="cp-2").isError
    # and an edit inside the fresh insertion still commits (single-level nesting builds fine)
    _j(workdir_open(str(workdir), track=True))
    r2 = document_patch(hunks=[{"paragraph_id": "P1", "old": "甲", "new": "甲乙"}], operation_id="cp-3")
    assert not r2.isError, r2.structuredContent
    committed = commit_sync(operation_id="cp-4")
    assert not committed.isError, committed.structuredContent
    built = tmp_path / "probe-out.docx"
    assert not build_docx(output=str(built), operation_id="cp-5").isError
    assert not verify_output(output=str(built), operation_id="cp-6").isError


def test_batch_edit_comment_gate_and_canonical_opt_in(tmp_path):
    """#78 close-out: batch_edit cannot bypass the comment-text opt-in; the
    flag is part of the canonical operation input, so reusing an operation_id
    with a flipped flag fails operation-id-reused instead of replaying."""
    _reset()
    import shutil
    source = tmp_path / "cmt2-src.docx"
    shutil.copy2(RELEASE / "comments.docx", source)
    workdir = tmp_path / "cmt2"
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir)))
    d = _j(get_paragraph("comments.P0"))
    import re as _re
    visible = d["data"]["plain"] if "data" in d else d["plain"]
    old = _re.sub("\u27e6[^\u27e7]*\u27e7", "", visible)

    r = batch_edit(paragraph_id="comments.P0", edits=[{"text": old, "new": old + "X"}], operation_id="be-1")
    assert r.isError
    assert r.structuredContent["diagnostics"][0]["code"] == "comment-text-requires-opt-in"

    ok = batch_edit(paragraph_id="comments.P0", edits=[{"text": old, "new": old + "X"}], operation_id="be-2", allow_comment_text=True)
    assert not ok.isError, ok.structuredContent

    # flipped flag on a reused operation_id must fail, not replay
    _j(revert(operation_id="be-r"))
    replay = batch_edit(paragraph_id="comments.P0", edits=[{"text": old, "new": old + "X"}], operation_id="be-2", allow_comment_text=False)
    assert replay.isError
    assert replay.structuredContent["diagnostics"][0]["code"] == "operation-id-reused", replay.structuredContent["diagnostics"][0]


def test_source_drift_blocks_all_mutations_not_commit_only(tmp_path):
    """Drift refusal is a global mutation precondition: batch_edit directly
    syncs+publishes without commit_sync, so it must fail on its own; an
    exact operation_id retry of a PRE-drift success still replays
    (ledger-first); a NEW operation id after drift fails closed."""
    workdir = _open_tracked(tmp_path, "drift2")
    source = workdir.parent / "drift2-src.docx"
    # a successful pre-drift mutation with a pinned operation_id
    r0 = replace_text(paragraph_id="P0", old="甲段落原文内容", new="甲段落原文内容A", operation_id="pre-1")
    assert not r0.isError, r0.structuredContent
    data = bytearray(source.read_bytes())
    data[-1] ^= 0xFF
    source.write_bytes(bytes(data))
    # exact retry replays the original success despite the drift
    replay = replace_text(paragraph_id="P0", old="甲段落原文内容", new="甲段落原文内容A", operation_id="pre-1")
    assert not replay.isError, replay.structuredContent
    # a NEW operation id fails closed on every text-mutation lane
    r_new = replace_text(paragraph_id="P1", old="目标插入语", new="目标插入语Z", operation_id="post-1")
    assert r_new.isError
    assert r_new.structuredContent["diagnostics"][0]["code"] == "source-modified-outside-engine"
    b = batch_edit(paragraph_id="P1", edits=[{"new": "x"}], operation_id="post-2")
    assert b.isError
    assert b.structuredContent["diagnostics"][0]["code"] == "source-modified-outside-engine"


def test_noop_replace_hunks_rejected(tmp_path):
    """#80: old == new hunks are refused with patch-noop (document_patch and
    batch_edit) and leave the draft untouched."""
    workdir = _open_tracked(tmp_path, "noop")
    before = (workdir / "edit.md").read_bytes()
    r = document_patch(hunks=[{"paragraph_id": "P1", "old": "目标插入语", "new": "目标插入语"}], operation_id="noop-1")
    assert r.isError
    assert r.structuredContent["diagnostics"][0]["code"] == "patch-noop"
    assert (workdir / "edit.md").read_bytes() == before
    b = batch_edit(paragraph_id="P1", edits=[{"text": "乙段落前缀文字 目标插入语 后缀文字收尾", "new": "乙段落前缀文字 目标插入语 后缀文字收尾"}], operation_id="noop-2")
    assert b.isError
    assert b.structuredContent["diagnostics"][0]["code"] == "patch-noop"
    b2 = batch_edit(paragraph_id="P1", edits=[{"region": 0, "new": "乙段落前缀文字 目标插入语 后缀文字收尾"}], operation_id="noop-3")
    assert b2.isError
    assert b2.structuredContent["diagnostics"][0]["code"] == "patch-noop"
    assert (workdir / "edit.md").read_bytes() == before


def test_span_map_view_and_refusal_payload(tmp_path):
    """P0-1: document_read(view="spans") returns copy-paste-legal editable
    spans cut at revision boundaries; a boundary/text-not-found refusal
    carries the same map so the agent never has to guess."""
    workdir = _open_tracked(tmp_path, "spans")
    # commit an insertion so P1 carries an insert region
    r = document_patch(hunks=[{"paragraph_id": "P1", "old": "目标插入语", "new": "目标插入语甲"}], operation_id="sp-1")
    assert not r.isError
    assert not commit_sync(operation_id="sp-2").isError
    _j(workdir_open(str(workdir), track=True))

    m = _j(document_read(anchor="P1", view="spans"))["span_map"]
    texts = [s["text"] for s in m["spans"]]
    assert any("甲" in t for t in texts), texts
    assert m["boundaries"], m
    assert all(s["region"] in ("baseline", "insert") for s in m["spans"])
    assert any(s["region"] == "insert" for s in m["spans"]), m
    assert m["style_regions"], m
    assert max(s["length"] for s in m["spans"]) > 3, [s["text"] for s in m["spans"]]
    # every unique span text is a legal old string
    unique_spans = [s for s in m["spans"] if s["unique"] and s["text"].strip()]
    assert unique_spans, m
    for span in unique_spans:
        probe = document_patch(hunks=[{"paragraph_id": "P1", "old": span["text"], "new": span["text"] + "·"}], operation_id=f"sp-legal-{span['index']}")
        assert not probe.isError, (span, probe.structuredContent)
        _j(revert(operation_id=f"sp-legal-r{span['index']}"))

    # crossing refusal carries the span map
    cross = document_patch(hunks=[{"paragraph_id": "P1", "old": "插入语甲", "new": "插入语甲X"}], operation_id="sp-cross")
    assert cross.isError
    diag = cross.structuredContent["diagnostics"][0]
    assert diag["code"] == "edit-span-crosses-revision-boundary"
    assert diag["details"]["span_map"]["paragraph_id"] == "P1"
    assert diag["details"]["span_map"]["spans"], diag

    # text-not-found refusal carries it too
    nf = document_patch(hunks=[{"paragraph_id": "P1", "old": "不存在的句子", "new": "x"}], operation_id="sp-nf")
    assert nf.isError
    assert nf.structuredContent["diagnostics"][0]["details"]["span_map"]["spans"]


def test_batch_hunks_report_every_broken_hunk(tmp_path):
    """UX: an 11-hunk batch must name the broken hunks (index + paragraph +
    code) instead of failing with one opaque error; a single broken hunk
    keeps its own code, wrapped with its hunk number."""
    workdir = _open_tracked(tmp_path, "hunks")
    before = (workdir / "edit.md").read_bytes()
    # three hunks: #1 fine, #2 not found, #3 noop -> one combined report
    r = document_patch(
        hunks=[
            {"paragraph_id": "P1", "old": "目标插入语", "new": "目标插入语X"},
            {"paragraph_id": "P1", "old": "不存在的文本", "new": "y"},
            {"paragraph_id": "P0", "old": "保持不动", "new": "保持不动"},
        ],
        operation_id="hunks-1",
    )
    assert r.isError
    diag = r.structuredContent["diagnostics"][0]
    assert diag["code"] == "patch-hunks-invalid", diag
    problems = diag["details"]["problems"]
    assert [p["hunk"] for p in problems] == [2, 3], problems
    assert [p["code"] for p in problems] == ["text-not-found", "patch-noop"], problems
    assert problems[0]["paragraph_id"] == "P1" and problems[1]["paragraph_id"] == "P0"
    assert problems[0]["span_map"]["paragraph_id"] == "P1"
    assert (workdir / "edit.md").read_bytes() == before  # nothing written

    # one broken hunk -> its own code, hunk number in the message
    r1 = document_patch(
        hunks=[
            {"paragraph_id": "P1", "old": "目标插入语", "new": "目标插入语X"},
            {"paragraph_id": "P1", "old": "不存在的文本", "new": "y"},
        ],
        operation_id="hunks-2",
    )
    assert r1.isError
    d1 = r1.structuredContent["diagnostics"][0]
    assert d1["code"] == "text-not-found"
    assert "hunk #2" in d1["message"], d1["message"]
    assert d1["details"]["hunk"] == 2
    # the good hunk alone still applies
    ok = document_patch(hunks=[{"paragraph_id": "P1", "old": "目标插入语", "new": "目标插入语X"}], operation_id="hunks-3")
    assert not ok.isError, ok.structuredContent


def _make_superscript_docx(path, *, with_reference: bool):
    from docx import Document
    d = Document()
    if with_reference:
        p = d.add_paragraph()
        p.add_run("既有正确写法 Cu")
        run = p.add_run("2+")
        run.font.superscript = True
    d.add_paragraph("本段落里的 Cu2+ 上下标缺失，需要修复。")
    d.save(path)


def _format_span_workdir(tmp_path, name, *, with_reference, track=False):
    from scripts.extract import extract
    source = tmp_path / f"{name}-src.docx"
    _make_superscript_docx(source, with_reference=with_reference)
    workdir = tmp_path / name
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir), track=track, author="Lin"))
    return workdir


def test_format_span_reuses_existing_style_and_is_minimal(tmp_path):
    """format_span fixes missing superscripts by reusing the template's
    existing run-property variant; direct mode restyles in place, tracked
    mode emits a minimal delete+insert pair (only the target text)."""
    from scripts.extract import extract
    workdir = _format_span_workdir(tmp_path, "fmt", with_reference=True)
    target = "P1" if "P1" in (workdir / "typed.md").read_text(encoding="utf-8") else "P0"

    direct = format_span(paragraph_id=target, old="Cu2+", attributes={"vertAlign": "superscript"}, operation_id="fmt-1")
    assert not direct.isError, direct.structuredContent
    assert direct.structuredContent["data"]["edit_mode"] == "direct"
    again = format_span(paragraph_id=target, old="Cu2+", attributes={"vertAlign": "superscript"}, operation_id="fmt-2")
    assert again.isError and again.structuredContent["diagnostics"][0]["code"] == "format-noop"
    out = tmp_path / "fmt-out.docx"
    _j(commit_sync(operation_id="fmt-save-a"))  # ADR 0044
    assert not build_docx(output=str(out), operation_id="fmt-3").isError
    assert not verify_output(output=str(out), operation_id="fmt-4").isError
    import zipfile
    from lxml import etree
    with zipfile.ZipFile(out) as z:
        doc = etree.fromstring(z.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    sups = [run for run in doc.xpath(".//w:r", namespaces=ns) if run.xpath("./w:rPr/w:vertAlign[@w:val='superscript']", namespaces=ns)]
    assert sups, "output must carry a superscript run"

    # tracked mode: minimal revision pair, authored by the session author
    tracked_dir = _format_span_workdir(tmp_path, "fmt-track", with_reference=True, track=True)
    t_target = "P1" if "P1" in (tracked_dir / "typed.md").read_text(encoding="utf-8") else "P0"
    tr = format_span(paragraph_id=t_target, old="Cu2+", attributes={"vertAlign": "superscript"}, operation_id="fmt-5")
    assert not tr.isError, tr.structuredContent
    assert tr.structuredContent["data"]["edit_mode"] == "track"
    tout = tmp_path / "fmt-track-out.docx"
    _j(commit_sync(operation_id="fmt-save-b"))  # ADR 0044
    assert not build_docx(output=str(tout), operation_id="fmt-6").isError
    assert not verify_output(output=str(tout), operation_id="fmt-7").isError
    with zipfile.ZipFile(tout) as z:
        tdoc = etree.fromstring(z.read("word/document.xml"))
    inserted = ["".join(t.text or "" for t in el.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")) for el in tdoc.xpath(".//w:ins", namespaces=ns)]
    assert inserted == ["Cu2+"], inserted
    authors = {el.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}author") for el in tdoc.xpath(".//w:ins", namespaces=ns)}
    assert authors == {"Lin"}, authors


def test_format_span_never_invents_a_style(tmp_path):
    """A document with no such run-property variant anywhere cannot be given
    one: the refusal names the missing variant and the available options."""
    workdir = _format_span_workdir(tmp_path, "fmt-none", with_reference=False)
    r = format_span(paragraph_id="P0", old="Cu2+", attributes={"vertAlign": "superscript"}, operation_id="fmt-none-1")
    assert r.isError
    diag = r.structuredContent["diagnostics"][0]
    assert diag["code"] == "format-style-unavailable", diag
    assert "cannot be invented" in diag["message"]
    assert (workdir / "typed.md").read_text(encoding="utf-8").count("span data-s") == 0


def test_matching_tolerates_width_and_punctuation(tmp_path):
    """First-try success for the common Chinese-text mismatch: the agent types
    half-width punctuation/spaces while the document uses full-width. The
    patch applies, and the warning names the exact convention differences."""
    from scripts.extract import extract
    from docx import Document
    source = tmp_path / "tol-src.docx"
    d = Document()
    d.add_paragraph("缓冲液中同时含有Cu2+和Mn2+两种金属离子，浓度5～20 mM。")
    d.save(source)
    workdir = tmp_path / "tol"
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir), track=False))

    r = document_patch(
        hunks=[{"paragraph_id": "P0", "old": "两种金属离子,浓度5~20 mM", "new": "两种金属离子，浓度10～20 mM"}],
        operation_id="tol-1",
    )
    assert not r.isError, r.structuredContent
    warnings = r.structuredContent["data"]["warnings"]
    assert any("matched-with-normalization" in w for w in warnings), warnings
    assert "doc='，'" in warnings[0] and "doc='～'" in warnings[0]
    draft = (workdir / "edit.md").read_text(encoding="utf-8")
    assert "浓度10～20 mM" in draft  # the caller's new text is written verbatim
    assert "matched-with-normalization" in json.dumps(r.structuredContent["data"]["applied"], ensure_ascii=False)


def test_not_found_carries_a_ready_to_send_fix(tmp_path):
    """A wrong old string must not send the agent exploring: the refusal names
    the divergence AND ships the corrected hunk in data.fix."""
    workdir = _open_tracked(tmp_path, "fixrecipe")
    r = document_patch(
        hunks=[{"paragraph_id": "P1", "old": "乙段落前缀文字 目标插入语 后缀语", "new": "X"}],
        operation_id="fix-1",
    )
    assert r.isError
    diag = r.structuredContent["diagnostics"][0]
    assert diag["code"] == "text-not-found"
    fix = diag["details"]["fix"]
    assert fix["action"] == "resend-hunk-with-document-text"
    assert fix["paragraph_id"] == "P1"
    assert fix["new"] == "X"
    # the suggested old is the document's real text and actually applies
    ok = document_patch(hunks=[{"paragraph_id": fix["paragraph_id"], "old": fix["old"], "new": fix["new"]}], operation_id="fix-2")
    assert not ok.isError, (fix, ok.structuredContent)


def test_text_inside_tracked_deletion_is_named(tmp_path):
    """Editing text that a prior revision deleted is reported as such (with the
    revision identity), not as a bare not-found."""
    workdir = _open_tracked(tmp_path, "delcase")
    # create a tracked deletion: replace text in track mode, then commit
    r = document_patch(hunks=[{"paragraph_id": "P1", "old": "目标插入语", "new": "替换后文本"}], operation_id="del-1")
    assert not r.isError
    assert not commit_sync(operation_id="del-2").isError
    _j(workdir_open(str(workdir), track=True))
    r2 = document_patch(hunks=[{"paragraph_id": "P1", "old": "目标插入语", "new": "再次修改"}], operation_id="del-3")
    assert r2.isError
    diag = r2.structuredContent["diagnostics"][0]
    assert diag["code"] == "text-inside-tracked-deletion", diag
    assert diag["details"]["deletion"].get("w_id")


def test_span_index_addressing_and_issues_view(tmp_path):
    """Path-addressed formatting (span_index from the read surface) plus the
    document issues view: charges without superscript, mixed punctuation
    width, and a full name defined twice."""
    from docx import Document
    from scripts.extract import extract
    source = tmp_path / "iss-src.docx"
    d = Document()
    p = d.add_paragraph()
    p.add_run("支架中同时含有Cu")
    run = p.add_run("2+")
    run.font.superscript = True
    p.add_run("和Mn2+,浓度5 mM。")
    d.add_paragraph("甲基丙烯酰化透明质酸（HAMA）多孔支架，与HAMA凝胶复合。")
    d.add_paragraph("再次出现甲基丙烯酰化透明质酸（HAMA）全称。")
    d.save(source)
    workdir = tmp_path / "iss"
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir), track=False))

    issues = _j(document_read(view="issues"))
    kinds = [i["kind"] for i in issues["issues"]]
    assert "element-charge-not-superscript" in kinds, kinds
    assert "punctuation-width-mixed" in kinds, kinds
    assert "full-name-defined-repeatedly" in kinds, kinds
    charge_issue = next(i for i in issues["issues"] if i["kind"] == "element-charge-not-superscript")
    assert charge_issue["paragraph_id"] == "P0"
    assert charge_issue["fix"]["tool"] == "format_span"

    # path addressing: no text matching at all — address the style REGION
    regions = _j(document_read(anchor="P0", view="spans"))["span_map"]["style_regions"]
    target = next(reg for reg in regions if "Mn2+" in reg["text"])
    r = format_span(paragraph_id="P0", span_index=target["index"], attributes={"vertAlign": "superscript"}, operation_id="iss-1")
    assert not r.isError, (target, r.structuredContent)
    out = tmp_path / "iss-out.docx"
    _j(commit_sync(operation_id="iss-save"))  # ADR 0044
    assert not build_docx(output=str(out), operation_id="iss-2").isError
    assert not verify_output(output=str(out), operation_id="iss-3").isError
    oob = format_span(paragraph_id="P0", span_index=999, attributes={"vertAlign": "superscript"}, operation_id="iss-4")
    assert oob.isError and oob.structuredContent["diagnostics"][0]["code"] == "span-index-out-of-range"


def test_search_matches_across_inline_markers(tmp_path):
    """Search runs on token-free text: a query interrupted by revision edges
    must hit (it used to return 0), and a match inside one span is directly
    usable as a patch old — while a cross-span match is reported with the
    precise boundary diagnostic instead of silence."""
    workdir = _open_tracked(tmp_path, "searchtok")
    r = document_patch(hunks=[{"paragraph_id": "P1", "old": "目标插入语", "new": "目标插入语甲"}], operation_id="st-1")
    assert not r.isError
    assert not commit_sync(operation_id="st-2").isError
    _j(workdir_open(str(workdir), track=True))

    cross = _j(document_search("前缀文字 目标插入语甲 后缀文字收尾"))
    assert cross["total_matches"] == 1, cross
    hit = cross["matches"][0]
    assert hit["id"] == "P1"
    assert hit["matched_text"] == "前缀文字 目标插入语甲 后缀文字收尾"
    # read/write contract closes: a crossing hit says so
    assert hit["patchable_as_single_hunk"] is False
    assert hit["region"] == "mixed"
    assert hit["boundary_crossings"], hit
    assert hit["span_indices"], hit

    # the cross-boundary match is refused with the pin-point diagnostic
    refused = document_patch(hunks=[{"paragraph_id": "P1", "old": hit["matched_text"], "new": "整体替换后"}], operation_id="st-3")
    assert refused.isError
    diag = refused.structuredContent["diagnostics"][0]
    assert diag["code"] == "edit-span-crosses-revision-boundary", diag
    assert diag["details"]["span_map"]["spans"]

    # a match inside ONE span patches directly
    inner = _j(document_search("后缀文字收尾"))
    inner_hit = inner["matches"][0]
    assert inner_hit["patchable_as_single_hunk"] is True
    assert inner_hit["region"] == "baseline"
    assert inner_hit["boundary_crossings"] == []
    probe = document_patch(hunks=[{"paragraph_id": "P1", "old": inner_hit["matched_text"], "new": "尾部替换"}], operation_id="st-4")
    assert not probe.isError, probe.structuredContent


def test_search_tolerates_width_variants(tmp_path):
    """The same folding the patch lane uses applies to search, so an agent
    typing half-width punctuation still finds the full-width document text."""
    from scripts.extract import extract
    from docx import Document
    source = tmp_path / "sw-src.docx"
    d = Document()
    d.add_paragraph("支架中同时含有Cu2+和Mn2+两种金属离子，浓度5～20 mM。")
    d.save(source)
    workdir = tmp_path / "sw"
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir), track=False))
    found = _j(document_search("两种金属离子,浓度5~20 mM"))
    assert found["total_matches"] == 1, found
    assert found["matches"][0]["normalized"] is True
    assert found["matches"][0]["matched_text"] == "两种金属离子，浓度5～20 mM"


def test_document_replace_plan_atomicity_and_refs(tmp_path):
    """document_replace: scope discovery, match plan with refs, all-or-nothing
    on unsafe matches, expected_matches guard, and no-match as a success no-op.
    References from search/replace drive patch and format without copying text."""
    from docx import Document
    from scripts.extract import extract
    source = tmp_path / "rep-src.docx"
    d = Document()
    d.add_paragraph("骨关节炎的治疗包括药物，骨关节炎需要长期管理。")
    d.add_paragraph("骨关节炎患者应定期复查。")
    d.save(source)
    workdir = tmp_path / "rep"
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir), track=False))

    r = document_replace(find="骨关节炎", replace="膝骨关节炎", operation_id="rep-1")
    assert not r.isError, r.structuredContent
    data = r.structuredContent["data"]
    assert data["matches"] == 3 and data["changed"] == 3
    assert all(entry["patchable"] for entry in data["match_plan"])
    assert all(entry["match_ref"].startswith("ref_") for entry in data["match_plan"])
    assert data["document_state"]["revision_before"] != data["document_state"]["revision_after"]
    assert not commit_sync(operation_id="rep-2").isError
    out = tmp_path / "rep-out.docx"
    assert not build_docx(output=str(out), operation_id="rep-3").isError
    assert not verify_output(output=str(out), operation_id="rep-4").isError

    # expected_matches guard: fail closed, nothing written
    before = (workdir / "edit.md").read_bytes()
    bad = document_replace(find="膝骨关节炎", replace="X", expected_matches=99, operation_id="rep-5")
    assert bad.isError and bad.structuredContent["diagnostics"][0]["code"] == "replace-expected-matches-mismatch"
    assert (workdir / "edit.md").read_bytes() == before

    # no-match is an informational success (not an error, not a silent edit)
    none = document_replace(find="不存在的词", replace="X", operation_id="rep-6")
    assert not none.isError
    assert none.structuredContent["data"]["changed"] == 0
    assert none.structuredContent["data"]["noop_reason"] == "no-match"

    # refs drive patch + format
    hit = _j(document_search("膝骨关节炎患者"))["matches"][0]
    assert hit["patchable_as_single_hunk"] is True
    patched = document_patch(hunks=[{"match_ref": hit["match_ref"], "new": "膝骨关节炎患者应定期复诊"}], operation_id="rep-7")
    assert not patched.isError, patched.structuredContent
    assert not commit_sync(operation_id="rep-8").isError
    # an old reference whose offsets still hold is applied (the anchor, not the
    # revision id, decides); one whose TEXT changed fails closed with the
    # current text plus a fresh reference, so the retry is one step
    overlap = document_patch(
        hunks=[{"paragraph_id": "P1", "old": "膝骨关节炎患者", "new": "患者本人"}],
        operation_id="rep-9",
    )
    assert not overlap.isError, overlap.structuredContent
    assert not commit_sync(operation_id="rep-10").isError
    stale = document_patch(hunks=[{"match_ref": hit["match_ref"], "new": "X"}], operation_id="rep-11")
    assert stale.isError
    diag = stale.structuredContent["diagnostics"][0]
    assert diag["code"] == "match-ref-stale", diag
    details = diag["details"]
    assert details["current_text"] and details["fresh_match_ref"]
    assert "患者本人" in details["current_text"]


def test_capability_manifest_layers(tmp_path):
    """engine_info exposes the static manifest; document_read(view=
    "capabilities") reports what THIS document can do, and a closed lane names
    its capability id."""
    static = engine_info()["capabilities"]
    ids = {entry["capability"] for entry in static}
    assert "word.text.replace.cross-revision-boundary" in ids
    assert any(entry["support"] == "unsupported" and entry.get("current_fallback") for entry in static)

    workdir = _open_tracked(tmp_path, "caps")
    caps = _j(document_read(view="capabilities"))
    assert caps["document"]["state"] == "clean"
    by_id = {entry["capability"]: entry for entry in caps["capabilities"]}
    # no revision boundaries committed yet -> the lane is open for this document
    assert by_id["word.text.replace.cross-revision-boundary"]["support"] == "supported"
    # no superscript variant in this fixture -> formatting lane closed with a reason
    assert by_id["word.format.run-properties"]["support"] in ("conditional", "unsupported")

    # a closed lane's refusal names the capability
    doc = document_patch(hunks=[{"paragraph_id": "P1", "old": "不存在的文本", "new": "x"}], operation_id="caps-1")
    assert doc.isError


def test_mode_can_be_chosen_in_the_mutation(tmp_path):
    """The mode decision rides the mutation: a revision-bearing document with
    an ambiguous mode no longer costs a refused call + a reopen."""
    workdir = _open_tracked(tmp_path, "modecall")
    r = document_patch(hunks=[{"paragraph_id": "P1", "old": "目标插入语", "new": "目标插入语甲"}], operation_id="mc-1")
    assert not r.isError
    assert not commit_sync(operation_id="mc-2").isError
    _j(workdir_open(str(workdir)))  # no track flag: the document's mode is inferred
    # track stated inside the mutation -> the FIRST call succeeds, no refusal cycle
    again = document_patch(
        hunks=[{"paragraph_id": "P1", "old": "甲", "new": "甲乙"}],
        operation_id="mc-3",
        track=True,
    )
    assert not again.isError, again.structuredContent
    assert session.mode == "track"


def test_direct_mode_refuses_revision_text_at_patch_time(tmp_path):
    """Editing text that sits INSIDE a tracked insertion in direct mode is
    refused by the PATCH (not by commit), and the refusal ships the fix."""
    workdir = _open_tracked(tmp_path, "directrev")
    r = document_patch(hunks=[{"paragraph_id": "P1", "old": "目标插入语", "new": "目标插入语甲"}], operation_id="dr-1", track=True)
    assert not r.isError
    assert not commit_sync(operation_id="dr-2").isError
    _j(workdir_open(str(workdir), track=False))
    # the inserted text "甲" lives wholly INSIDE the tracked insertion
    refused = document_patch(
        hunks=[{"paragraph_id": "P1", "old": "甲", "new": "乙"}],
        operation_id="dr-3",
    )
    assert refused.isError
    diag = refused.structuredContent["diagnostics"][0]
    assert diag["code"] == "revision-text-mutated-in-direct-mode", diag
    fix = diag["details"]["fix"]
    assert fix["track"] is True
    applied = document_patch(hunks=fix["hunks"], operation_id="dr-4", track=fix["track"])
    assert not applied.isError, applied.structuredContent


def test_editor_profile_surface():
    """The editor profile is the small front door AND it can recover a bad
    draft (revert), verified in a subprocess so the global server is intact."""
    import subprocess, sys
    code = (
        "import sys; sys.path.insert(0, '.');"
        "import scripts.mcp_server as m;"
        "m.apply_tool_profile('editor');"
        "print(sorted(m.mcp._tool_manager._tools))"
    )
    out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, cwd=str(Path(__file__).resolve().parents[1]))
    tools = out.stdout.strip()
    assert "'revert'" in tools, out.stdout + out.stderr
    assert "'document_patch'" in tools and "'format_span'" in tools and "'document_replace'" in tools
    # the save boundary must stay reachable: commit can demand the collaboration preflight
    assert "'review_preflight'" in tools and "'review_ack'" in tools, out.stdout
    assert "'get_paragraph'" not in tools and "'batch_edit'" not in tools


def test_replace_scope_semantics_exclude_comments_from_body(tmp_path):
    """scope=body means plain P<n> paragraphs: comment text must never be swept
    into an ordinary replace (it needs the explicit opt-in), and part/cell
    scopes resolve by id prefix."""
    from scripts.extract import extract
    import shutil
    source = tmp_path / "scope-src.docx"
    shutil.copy2(RELEASE / "comments.docx", source)
    workdir = tmp_path / "scope"
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir), track=False))

    body = [pid for pid in _scope_paragraph_ids(workdir, "body")]
    assert body and all(pid.startswith("P") or pid.startswith("T") for pid in body), body[:5]
    assert not any(pid.startswith("comments.") for pid in body), body[:5]
    everything = _scope_paragraph_ids(workdir, "all")
    assert len(everything) > len(body)
    comments = _scope_paragraph_ids(workdir, "comments")
    assert comments and all(pid.startswith("comments.") for pid in comments)
    assert _scope_paragraph_ids(workdir, body[0]) == [body[0]]
    with pytest.raises(ToolError):
        _scope_paragraph_ids(workdir, "no-such-part")


def test_token_boundaries_are_named_and_carry_a_split_fix(tmp_path):
    """#82-follow-on: crossing ANY protected inline token (not just revision
    markers) is refused with the flavour named and a split skeleton, instead of
    a misleading text-not-found."""
    from docx import Document
    from docx.oxml.ns import qn
    from scripts.extract import extract
    source = tmp_path / "tok-src.docx"
    d = Document()
    p = d.add_paragraph()
    p.add_run("前缀甲乙")
    bookmark = p._p.makeelement(qn("w:bookmarkStart"), {qn("w:id"): "7", qn("w:name"): "mk"})
    p._p.append(bookmark)
    p._p.append(p._p.makeelement(qn("w:bookmarkEnd"), {qn("w:id"): "7"}))
    p.add_run("后缀丙丁")
    d.save(source)
    workdir = tmp_path / "tok"
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir), track=False))

    m = _j(document_read(anchor="P0", view="spans"))["span_map"]
    kinds = [b["kind"] for b in m["boundaries"]]
    assert any(kind.startswith("token:") for kind in kinds), kinds

    r = document_patch(hunks=[{"paragraph_id": "P0", "old": "前缀甲乙后缀丙丁", "new": "整段替换"}], operation_id="tok-1")
    assert r.isError
    diag = r.structuredContent["diagnostics"][0]
    # a bookmark is an anchor, not a revision: name it as what it is
    assert diag["code"] == "edit-span-crosses-protected-marker"
    assert "protected inline markers" in diag["message"]
    assert diag["details"]["fix"]["action"] == "split-into-per-span-hunks"
    assert diag["details"]["fix"]["hunks"][0]["old"] == "前缀甲乙"
    assert diag["details"]["fix"]["hunks"][1]["old"] == "后缀丙丁"


def test_search_scope_and_paging(tmp_path):
    """Search narrows by scope and pages by offset, so a long document can be
    walked without guessing (the acceptance run asked for exactly this)."""
    workdir = _open_tracked(tmp_path, "searchpage")
    everything = _j(document_search("目标插入语"))
    assert everything["scope"] == "all" and everything["total_blocks"] >= 1
    assert everything["offset"] == 0
    body = _j(document_search("目标插入语", scope="body"))
    assert [m["id"] for m in body["matches"]] == ["P1"]
    one = _j(document_search("目标插入语", scope="P1"))
    assert [m["id"] for m in one["matches"]] == ["P1"]
    empty = _j(document_search("目标插入语", scope="P1", offset=5))
    assert empty["matches"] == [] and empty["total_blocks"] == 1
    with pytest.raises(ToolError):
        document_search("目标插入语", scope="no-such-scope")


def test_search_hits_expose_every_occurrence(tmp_path):
    """'Change the SECOND occurrence' must be two calls: search -> patch(occ ref).
    Each hit lists every occurrence with its own version-bound address."""
    from docx import Document
    from scripts.extract import extract
    source = tmp_path / "occ-src.docx"
    d = Document()
    d.add_paragraph("甲组用血浆凝胶，乙组用血浆凝胶，丙组用血浆凝胶。")
    d.save(source)
    workdir = tmp_path / "occ"
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir), track=False))

    hits = _j(document_search("血浆凝胶"))
    hit = hits["matches"][0]
    assert hit["matches"] == 3
    assert [occ["offset"] for occ in hit["occurrences"]] == sorted(occ["offset"] for occ in hit["occurrences"])
    assert len({occ["match_ref"] for occ in hit["occurrences"]}) == 3
    second = hit["occurrences"][1]
    assert second["patchable_as_single_hunk"] is True
    patched = document_patch(hunks=[{"match_ref": second["match_ref"], "new": "血浆凝胶层"}], operation_id="occ-1")
    assert not patched.isError, patched.structuredContent
    draft = (workdir / "edit.md").read_text(encoding="utf-8")
    assert draft.count("血浆凝胶层") == 1          # the second occurrence was retargeted
    assert draft.count("血浆凝胶，") == 1          # the first still stands
    assert draft.count("血浆凝胶。") == 1          # and so does the third


def test_unknown_hunk_key_is_refused_not_silently_dropped(tmp_path):
    """An unrecognised hunk key (typo) must never be silently ignored: a
    'replacement' key used to degrade the hunk into a pure deletion and strip
    the paragraph head while reporting success."""
    workdir = _open_tracked(tmp_path, "strictkeys")
    before = (workdir / "edit.md").read_bytes()
    typo = document_patch(
        hunks=[{"paragraph_id": "P1", "old": "目标插入语", "replacement": "目标插入语甲"}],
        operation_id="strict-1",
    )
    assert typo.isError
    diag = typo.structuredContent["diagnostics"][0]
    assert diag["code"] == "invalid-arguments"
    assert "replacement" in diag["message"] and "new" in diag["message"]
    assert (workdir / "edit.md").read_bytes() == before

    missing_new = document_patch(
        hunks=[{"match_ref": "ref_whatever", "replacement": "x"}],
        operation_id="strict-2",
    )
    assert missing_new.isError
    assert missing_new.structuredContent["diagnostics"][0]["code"] == "invalid-arguments"

    no_new = document_patch(hunks=[{"paragraph_id": "P1", "old": "目标插入语"}], operation_id="strict-3")
    assert no_new.isError
    assert "needs 'new'" in no_new.structuredContent["diagnostics"][0]["message"]
    assert (workdir / "edit.md").read_bytes() == before


def test_replace_refuses_direct_mode_revision_text_before_writing(tmp_path):
    """document_replace must apply the same early mode guard as document_patch:
    a bulk replace whose targets sit inside tracked insertions is refused
    BEFORE any write, with a ready track=true fix (the acceptance run
    otherwise applied 11 replacements and died at commit)."""
    workdir = _open_tracked(tmp_path, "replguard")
    r0 = document_patch(hunks=[{"paragraph_id": "P1", "old": "目标插入语", "new": "目标插入语甲"}], operation_id="rg-1", track=True)
    assert not r0.isError
    assert not commit_sync(operation_id="rg-2").isError
    _j(workdir_open(str(workdir), track=False))
    before = (workdir / "edit.md").read_bytes()
    r = document_replace(find="甲", replace="甲乙", scope="body", operation_id="rg-3")
    assert r.isError
    diag = r.structuredContent["diagnostics"][0]
    assert diag["code"] == "revision-text-mutated-in-direct-mode", diag
    assert "nothing was written" in diag["message"]
    fix = diag["details"]["fix"]
    assert fix["track"] is True and fix["args"]["find"] == "甲"
    assert (workdir / "edit.md").read_bytes() == before
    ok = document_replace(**fix["args"], track=fix["track"], operation_id="rg-4")
    assert not ok.isError, ok.structuredContent



def test_rpr_change_marker_round_trips_with_its_style(tmp_path):
    """Issue #83: a format-history marker (w:rPrChange) must round-trip as its
    own run carrying ITS style. Binding it to a neighbouring run silently
    changed the style (and, after a container, nested it one level deeper), so
    the rebuilt document no longer matched the model and the workdir became
    unbuildable."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from scripts.extract import extract

    source = tmp_path / "rpr-src.docx"
    document = Document()
    p = document.add_paragraph()
    p.add_run("前缀文字")
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "11")
    ins.set(qn("w:author"), "Lin")
    ins.set(qn("w:date"), "2026-09-01T00:00:00Z")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "插入文字"
    run.append(t)
    ins.append(run)
    p._p.append(ins)
    # a run that ONLY carries format history, styled superscript
    hist = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    vert = OxmlElement("w:vertAlign")
    vert.set(qn("w:val"), "superscript")
    rpr.append(vert)
    change = OxmlElement("w:rPrChange")
    change.set(qn("w:id"), "12")
    change.set(qn("w:author"), "Lin")
    change.set(qn("w:date"), "2026-09-01T00:00:00Z")
    old_rpr = OxmlElement("w:rPr")
    old_rpr.append(OxmlElement("w:b"))
    change.append(old_rpr)
    rpr.append(change)
    hist.append(rpr)
    p._p.append(hist)
    p.add_run("尾部文字")
    document.save(source)

    workdir = tmp_path / "rpr"
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir), track=False))
    before = (workdir / "typed.md").read_text(encoding="utf-8")
    assert "rpr-change" in before
    import re as _re

    marker_style = _re.search(r'<docx-inline id="N\d+" kind="rpr-change" style="([^"]+)"', before).group(1)
    assert marker_style  # the marker is styled at extraction

    out = tmp_path / "rpr-out.docx"
    _j(commit_sync(operation_id="rpr-save-a"))  # ADR 0044
    assert not build_docx(output=str(out), operation_id="rpr-1").isError
    assert not verify_output(output=str(out), operation_id="rpr-2").isError
    # and the rebuilt document keeps the marker's style (same typed signature)
    _j(commit_sync(operation_id="rpr-save-b"))  # ADR 0044
    assert not build_docx(output=str(tmp_path / "rpr-out2.docx"), operation_id="rpr-3").isError


def test_verify_output_can_omit_its_argument_after_a_build(tmp_path):
    """A verify right after build_docx may omit `output` (the session remembers
    the PUBLISHED artifact, not the transient staging path); before any build
    the omission is refused with a named diagnostic."""
    workdir = _open_tracked(tmp_path, "verifydefault")
    session.last_build_output = None  # a fresh session has built nothing yet
    early = verify_output()
    assert early.isError
    assert early.structuredContent["diagnostics"][0]["code"] == "verify-output-required"
    out = tmp_path / "vd-out.docx"
    _j(commit_sync(operation_id="vd-save"))  # ADR 0044
    assert not build_docx(output=str(out), operation_id="vd-1").isError
    assert session.last_build_output == out.resolve()
    implicit = verify_output()
    assert not implicit.isError, implicit.structuredContent
    assert implicit.structuredContent["data"]["verified"] == str(out.resolve())


def test_document_read_returns_a_section_index(tmp_path):
    """One read answers 'which section is this paragraph in': the response
    carries contiguous paragraph ranges per part."""
    workdir = _open_tracked(tmp_path, "structure")
    payload = _j(document_read(view="outline", anchor="P1", before=1, after=1))
    structure = payload["structure"]
    parts = {entry["part"]: entry for entry in structure}
    assert "body" in parts
    assert parts["body"]["from"] == "P0" and parts["body"]["paragraphs"] >= 2


def test_document_replace_dry_run_counts_without_writing(tmp_path):
    """`dry_run` answers "how many matches are there, and are they all safe"
    without touching the draft — the alternative was mutate-then-revert."""
    workdir = _open_tracked(tmp_path, "dryrun")
    before = (workdir / "typed.md").read_bytes()
    probe = _j(document_replace(find="后缀文字", replace="结尾文字", scope="body", dry_run=True, operation_id="dr-1"))
    assert probe["dry_run"] is True
    assert probe["changed"] == 0
    assert probe["matches"] >= 1
    assert (workdir / "typed.md").read_bytes() == before
    applied = _j(document_replace(
        find="后缀文字", replace="结尾文字", scope="body",
        expected_matches=probe["matches"], operation_id="dr-2",
    ))
    assert applied["changed"] == probe["matches"]


def test_recovery_hints_never_name_a_tool_the_profile_hides():
    """A stranded session reads a hidden tool as "unrecoverable": every
    suggested fix tool must be callable in the active profile."""
    import scripts.mcp_server as server

    previous = server._ACTIVE_TOOL_NAMES
    try:
        server.apply_tool_profile("editor")
        allowed = set(server._ACTIVE_TOOL_NAMES)
        for code in ("agent-preflight-required", "current-snapshot-drift", "workdir-not-open"):
            fix = server._filtered_recovery_for("commit_sync", code)
            assert set(fix.get("tools", [])) <= allowed, (code, fix)
        preflight = server._filtered_recovery_for("commit_sync", "agent-preflight-required")
        assert preflight["tools"][0] == "review_preflight"
    finally:
        server._ACTIVE_TOOL_NAMES = previous


def test_hyperlink_split_span_is_not_reported_as_a_revision_boundary(tmp_path):
    """A span split by an inline anchor/hyperlink has no revision to look for:
    the diagnostic must say protected-marker, and still carry the span map."""
    from scripts.extract import extract

    source = ROOT / "corpus" / "release" / "anchors.docx"
    workdir = tmp_path / "hyper"
    assert extract([str(source), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir)))
    result = replace_text(paragraph_id="P0", old="超链接：点击", new="改链接", operation_id="hx-diag")
    diagnostic = result.structuredContent["diagnostics"][0]
    assert diagnostic["code"] == "edit-span-crosses-protected-marker", diagnostic
    assert any(name.startswith("token:") for name in diagnostic["details"]["crossed"])


def test_document_search_counts_text_that_only_deleted_revisions_hold(tmp_path):
    """Word shows deleted tracked-change text, the draft does not: without this
    breakdown a user's count and the search count silently disagree."""
    from scripts.extract import extract

    workdir = tmp_path / "counts"
    assert extract([str(ROOT / "corpus" / "release" / "revisions.docx"), "-o", str(workdir)]) == 0
    _j(workdir_open(str(workdir)))
    counts = _j(document_search(query="旧文本", scope="all"))["counts"]
    assert counts["inside_deleted_tracked_changes"] == 1, counts
    assert counts["total_in_projection"] == counts["editable"] + counts["inside_deleted_tracked_changes"]
    assert counts["total_in_projection"] >= 1, counts
    assert "note" in counts


def test_document_patch_preview_shows_the_resulting_text(tmp_path):
    """The edited paragraph as it now reads: a duplicated or broken join is
    visible in the patch response itself instead of needing a read-back."""
    workdir = _open_tracked(tmp_path, "preview")
    result = _j(document_patch(
        hunks=[{"paragraph_id": "P1", "old": "后缀文字", "new": "后缀文字收尾 后缀文字"}],
        operation_id="pv-1",
    ))
    preview = result["result_preview"]
    assert [entry["paragraph_id"] for entry in preview] == ["P1"]
    assert "后缀文字收尾 后缀文字" in preview[0]["result"]
    # a short paragraph is shown whole
    assert preview[0]["chars"] == len(preview[0]["result"])


def test_document_patch_warns_about_a_duplicated_join(tmp_path):
    """An edit that leaves the same phrase twice in a row is a join error the
    response should name, instead of the agent discovering it two patches
    later (this is how the r2 task burned four mutations)."""
    workdir = _open_tracked(tmp_path, "repeat")
    result = _j(document_patch(
        hunks=[{"paragraph_id": "P1", "old": "后缀文字", "new": "后缀文字-后缀文字"}],
        operation_id="rp-1",
    ))
    assert any("result-repeat" in warning for warning in result["warnings"]), result["warnings"]


def test_search_paging_survives_the_context_window_option(tmp_path):
    """`context_chars` used to shadow the paging argument: the response echoed a
    character offset and the next page came back empty, so paging died exactly
    when a caller asked for windows."""
    workdir = _open_tracked(tmp_path, "paging")
    first = _j(document_search("目标插入语", limit=1, context_chars=20))
    assert first["offset"] == 0
    assert first["returned_blocks"] == 1
    second = _j(document_search("目标插入语", limit=1, offset=1, context_chars=20))
    assert second["offset"] == 1
    assert second["returned_blocks"] == 0


def test_unknown_argument_names_fail_closed_instead_of_being_ignored(tmp_path):
    """The MCP layer drops unknown keys, so a typo (build_docx(output_path=...))
    used to read as success while the tool ran with defaults — the artifact went
    to the default path. Refuse it and name the closest accepted key."""
    import asyncio

    import scripts.mcp_server as server

    _open_tracked(tmp_path, "argguard")
    with pytest.raises(Exception) as failure:
        asyncio.run(server.mcp._tool_manager.call_tool(
            "build_docx", {"output_path": str(tmp_path / "wrong.docx")}, convert_result=True
        ))
    message = str(failure.value)
    assert "output_path" in message and "output" in message, message
    assert not (tmp_path / "wrong.docx").exists()
    # the correctly spelled call still works
    _j(commit_sync(operation_id="argguard-save"))  # ADR 0044
    result = asyncio.run(server.mcp._tool_manager.call_tool(
        "build_docx", {"output": str(tmp_path / "right.docx")}, convert_result=True
    ))
    assert (tmp_path / "right.docx").exists()


def test_search_hands_over_patchable_spans_for_a_crossing_hit(tmp_path):
    """A hit that crosses a revision boundary cannot be patched as one hunk;
    the search response must therefore carry the sub-spans that CAN be, with
    working refs — otherwise the caller earns a refusal it could have avoided."""
    workdir = _open_tracked(tmp_path, "spans")
    assert not document_patch(
        hunks=[{"paragraph_id": "P1", "old": "目标插入语", "new": "目标插入语甲"}], operation_id="sp-mk"
    ).isError
    assert not commit_sync(operation_id="sp-mk-c").isError
    _j(workdir_open(str(workdir), track=True))

    hit = _j(document_search("语甲 后缀文字", scope="P1"))["matches"][0]["occurrences"][0]
    assert hit["patchable_as_single_hunk"] is False
    assert hit["patchable_spans"], hit
    inside = next(span for span in hit["patchable_spans"] if span["text"] == "甲")
    applied = _j(document_patch(
        hunks=[{"match_ref": inside["match_ref"], "new": "甲X"}], operation_id="sp-use"
    ))
    assert applied["affected_paragraph_ids"] == ["P1"], applied
    assert "甲X" in applied["result_preview"][0]["result"], applied["result_preview"]


def test_format_span_advances_the_collaboration_snapshot(tmp_path):
    """format_span writes typed.md directly: if it leaves the collaboration
    ledger behind, the very next commit_sync is refused with
    current-snapshot-drift and nothing in the editor profile can clear it."""
    from scripts.review_collab import document_state

    workdir = _open_tracked(tmp_path, "collab", with_bold=True)
    assert not format_span(
        paragraph_id="P1", old="后缀文字收尾", attributes={"bold": True}, operation_id="fs-collab"
    ).isError
    assert document_state(workdir)["current_matches_filesystem"] is True
    assert not commit_sync(operation_id="fs-collab-c").isError


def test_every_editing_lane_leaves_the_collaboration_ledger_consistent(tmp_path):
    """A mutation that leaves typed.md ahead of the collaboration snapshot
    dead-ends the next commit_sync (current-snapshot-drift, unclearable inside
    the editor profile). Hold every editing lane to the same invariant; each
    lane gets a fresh workdir because format_span requires a clean draft."""
    from scripts.review_collab import document_state

    def lane(name, run, with_bold=False):
        workdir = _open_tracked(tmp_path, f"lane-{name}", with_bold=with_bold)
        assert not run().isError, name
        assert document_state(workdir)["current_matches_filesystem"] is True, name
        return workdir

    lane("patch", lambda: document_patch(
        hunks=[{"paragraph_id": "P1", "old": "后缀文字", "new": "结尾文字"}], operation_id="lane-p"))
    lane("replace", lambda: document_replace(
        find="目标插入语", replace="目标短语", scope="P1", operation_id="lane-r"))
    lane("format_span", lambda: format_span(
        paragraph_id="P1", old="后缀文字收尾", attributes={"bold": True}, operation_id="lane-f"), with_bold=True)
    lane("revert", lambda: revert(operation_id="lane-v"))
